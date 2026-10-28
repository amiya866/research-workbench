# -*- coding: utf-8 -*-
"""由本地缓存、行情和锌企财报生成与锡镍同结构的锌研究周报。"""
import csv
import datetime
import json
import os
import re
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

import build_nickel_weekly as shared

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'zinc-weekly'
CACHE = Path(r'D:/Kimi/金属总网/网站构建/db/数据缓存_v1.db')
SHFE_CSV = Path(r'D:/Kimi/有色截面配置_2026/_价格缓存.csv')
LME_CSV = Path(r'D:/Kimi/有色截面配置_2026/_LME价格缓存.csv')
FINANCIAL = Path(r'D:/Kimi/金属总网/网站构建/财报汇总/锌/全球锌企季度产量梳理.xlsx')
META = BASE / 'data' / 'zinc_weekly_freshness.json'

SECTIONS = [
    ('sumsec', '本周总结', '总览', []),
    ('price', '价格、升贴水与进口盈亏', '宏观与策略', ['overview_shfe', 'overview_lme', 'ID02038785', 'a10097491', 'FU00024317', 'ID01030236']),
    ('supply_mine', '锌矿与进口', '重点关注/供给', ['ID01320080', 'ID01320188', 'ID00188330', 'a10001843', 'a10020591', 'a10020592', 'a10020599', 'a10020600']),
    ('supply_smelt', '精炼锌', '重点关注/供给', ['ID01510883', 'a10018029', 'a10018062']),
    ('inventory', '显性库存', '重点关注/供给', ['overview_global_stock', 'FU00016165', 'ID00188293', 'a10134450', 'ID00188329']),
    ('demand', '镀锌/压铸/终端', '重点关注/需求', ['ID00187499', 'a10097188', 'a10018028', 'a12803961', 'a10001884', 'a10001831', 'a12756052', 'GM0000033132']),
    ('cost', '成本与平衡', '平衡表', ['ID01241108']),
    ('framework', '产业框架', '总览', []),
    ('technical', '技术分析', '宏观与策略', ['overview_shfe']),
    ('sources', '数据缺口', '总览', []),
]
NAMES = {key: title for key, title in []}
UNITS = {'overview_shfe': '元/吨', 'overview_lme': '美元/吨', 'overview_global_stock': '吨'}
KPI = [('shfe', '沪锌主连', '元/吨'), ('lme', 'LME锌', '美元/吨'), ('lme03', 'LME 0-3', '美元/吨'),
       ('global_stock', 'LME+SHFE库存', '吨'), ('social_stock', '锌锭社库', '万吨'), ('tc', '锌精矿TC', '元/吨')]


def series(conn, ind_id, name, unit):
    rows = conn.execute('select date,value from series_points where ind_id=? order by date', (ind_id,)).fetchall()
    return {'labels': [r[0] for r in rows], 'datasets': [{'label': name, 'data': [r[1] for r in rows],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}, unit


def csv_chart(path, column, name):
    labels, values = [], []
    with open(path, encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            value = row.get(column)
            if value not in (None, ''):
                labels.append(row[''])
                values.append(float(value))
    return {'labels': labels, 'datasets': [{'label': name, 'data': values, 'borderColor': '#1F4E79',
            'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def last(chart):
    point = shared.last_point(chart)
    return [point[0], point[1]] if point else [None, None]


def global_stock(lme, shfe):
    dates = sorted(set(lme['labels']) | set(shfe['labels']))
    lv, sv = lme['datasets'][0]['data'], shfe['datasets'][0]['data']
    lm = dict(zip(lme['labels'], lv))
    sm = dict(zip(shfe['labels'], sv))
    values, last_l, last_s = [], None, None
    for date in dates:
        last_l = lm.get(date, last_l)
        last_s = sm.get(date, last_s)
        values.append(last_l + last_s if last_l is not None and last_s is not None else None)
    return {'labels': dates, 'datasets': [{'label': 'LME+SHFE库存', 'data': values,
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def rsi(values, period=14):
    if len(values) <= period:
        return None
    changes = [b - a for a, b in zip(values[-period - 1:-1], values[-period:])]
    gains = sum(max(x, 0) for x in changes) / period
    losses = sum(max(-x, 0) for x in changes) / period
    return 100 if not losses else 100 - 100 / (1 + gains / losses)


def ema(values, period):
    alpha = 2 / (period + 1)
    out = []
    for i, value in enumerate(values):
        out.append(value if i == 0 else alpha * value + (1 - alpha) * out[-1])
    return out


def update_daily(summary):
    path = BASE / 'reports' / 'daily_latest.html'
    if not path.exists():
        return
    page = path.read_text(encoding='utf-8')
    begin, end = '<!-- ZINC_WEEKLY_CARD_BEGIN -->', '<!-- ZINC_WEEKLY_CARD_END -->'
    metrics = ''.join('<li><b>%s</b>　%s</li>' % (shared.esc(m['name']), shared.esc(m['value'])) for m in summary['metrics'][:4])
    card = '''%s
<div class="card" id="zincWeeklyCard" style="border-left:3px solid var(--gold);margin:14px 0;">
  <div style="display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap;align-items:center;">
    <h3 style="margin:0;color:var(--gold);">研究周报 · 锌</h3>
    <a href="../zinc-weekly/#sumsec" style="font-weight:700;">打开完整锌周报 ↗</a>
  </div>
  <p style="margin-top:8px;">%s</p>
  <ul style="margin:8px 0 0 18px;font-size:12px;color:var(--dim);">%s</ul>
  <p style="font-size:11px;color:var(--dim);margin-top:8px;">数据截至 %s · 本地生成，随研究站管线自动更新 · <a href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">锌企财报</a></p>
</div>
%s''' % (begin, shared.esc(summary['headline']), metrics, shared.esc(summary['asof']), end)
    pattern = re.escape(begin) + r'.*?' + re.escape(end)
    if re.search(pattern, page, re.S):
        page = re.sub(pattern, lambda _m: card, page, count=1, flags=re.S)
    else:
        anchor = '<!-- NICKEL_WEEKLY_CARD_END -->'
        if anchor not in page:
            anchor = '<section class="dp on" id="dp1">'
        page = page.replace(anchor, anchor + '\n' + card, 1)
    path.write_text(page, encoding='utf-8')


def main():
    conn = sqlite3.connect(f'file:{CACHE.as_posix()}?mode=ro', uri=True)
    charts, source_meta, names, units = {}, {}, dict(NAMES), dict(UNITS)
    for ind_id, name, unit in conn.execute('select ind_id,name,unit from series_meta where comm="ZN"'):
        charts[ind_id], units[ind_id] = series(conn, ind_id, name, unit)
        source_meta[ind_id] = {'type': 'zhiji', 'indicator': ind_id, 'unit': unit}
        names[ind_id] = name
    conn.close()
    charts['overview_shfe'] = csv_chart(SHFE_CSV, 'ZN', '沪锌主连')
    charts['overview_lme'] = csv_chart(LME_CSV, 'ZN', 'LME锌')
    charts['overview_global_stock'] = global_stock(charts['a10134450'], charts['ID00188293'])
    for key, label in [('overview_shfe', 'guan+本地日K'), ('overview_lme', '新浪外盘+本地缓存'), ('overview_global_stock', 'LME+SHFE合计')]:
        source_meta[key] = {'type': 'derived', 'formula': label}
        names[key] = {'overview_shfe': '沪锌主连', 'overview_lme': 'LME锌', 'overview_global_stock': '全球显性库存'}[key]
    shared.NAMES = names
    shared.SECTIONS = SECTIONS
    built = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
    data = {'charts': charts, 'sourceMeta': {'charts': source_meta}, 'meta': {'builtAt': built, 'warnings': [], 'emptyCharts': []}}
    freshness = shared.build_freshness(data)
    META.write_text(json.dumps(freshness, ensure_ascii=False, indent=2), encoding='utf-8')
    latest = {
        'shfe': last(charts['overview_shfe']), 'lme': last(charts['overview_lme']),
        'lme03': last(charts['a10097491']), 'global_stock': last(charts['overview_global_stock']),
        'social_stock': last(charts['ID00188329']), 'tc': last(charts['ID01320080']),
    }
    asof = max(value[0] for value in latest.values() if value[0])
    shfe_values = [v for v in charts['overview_shfe']['datasets'][0]['data'] if v is not None]
    chg5 = shfe_values[-1] / shfe_values[-6] - 1 if len(shfe_values) > 6 else None
    stock_values = [v for v in charts['a10134450']['datasets'][0]['data'] if v is not None]
    stock_chg = stock_values[-1] / stock_values[-6] - 1 if len(stock_values) > 6 else None
    ma5 = sum(shfe_values[-5:]) / 5
    ma20 = sum(shfe_values[-20:]) / 20
    current_rsi = rsi(shfe_values)
    macd = [a - b for a, b in zip(ema(shfe_values, 12), ema(shfe_values, 26))]
    signal = ema(macd, 9)
    price_word = '偏强' if (chg5 or 0) > 0.01 else ('偏弱' if (chg5 or 0) < -0.01 else '震荡')
    stock_word = '累库' if (stock_chg or 0) > 0 else '去库'
    headline = ('本周定调：沪锌%s与LME%s并存 —— 沪锌近5个交易日 %s，LME库存近5个交易日 %s；'
                '锌锭社库 %s 万吨（%s），昆明50%% TC %s 元/吨（%s）。'
                '近端重点跟踪LME仓单、国内社库、TC与进口盈亏能否同向验证；⚙️规则摘要，不构成投资建议。') % (
        price_word, stock_word, shared.pct(chg5), shared.pct(stock_chg), shared.fmt(latest['social_stock'][1]), latest['social_stock'][0],
        shared.fmt(latest['tc'][1]), latest['tc'][0])
    metrics = [{'name': name, 'value': '%s %s（%s）' % (shared.fmt(latest[key][1]), unit, latest[key][0])} for key, name, unit in KPI]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'lib').mkdir(exist_ok=True)
    for name in ('research.css', 'weekly_mobile.css', 'weekly_polish.css', 'weekly_mobile.js'):
        shutil.copyfile(BASE / 'nickel-weekly' / name, OUT / name)
    shutil.copyfile(BASE / 'nickel-weekly' / 'lib' / 'chart.min.js', OUT / 'lib' / 'chart.min.js')
    shutil.copyfile(BASE / 'scripts' / 'zinc_weekly_research.js', OUT / 'research.js')
    base_css = re.search(r'<style>(.*?)</style>', (BASE / 'nickel-weekly' / 'index.html').read_text(encoding='utf-8'), re.S).group(1)

    parts = []
    for sid, title, trail, keys in SECTIONS:
        parts.append('<section class="report-page" data-page="%s" data-trail="%s"><div class="page-heading"><h2>%s</h2><span>%s</span></div><div class="page-body">' % (
            sid, shared.esc(trail), shared.esc(title), shared.esc(trail.replace('/', ' / ').upper())))
        if sid == 'sumsec':
            parts.append('<div class="tone"><b>%s</b></div><div class="roll">策略每日滚动更新：周报为周五存档，最新决策见日报页。使用左侧目录切换研究内容。</div>' % shared.esc(headline))
            parts.append('<details class="fold" id="q4outlook"><summary>后续观察与证伪条件（点击展开/折叠）</summary><div class="foldbody"><ul class="sq"><li>LME与SHFE库存：仓单下降需连续验证，不以单日变化外推趋势。</li><li>矿端与冶炼：TC、进口矿、精炼锌产量共同判断供应，不将单一减产传闻直接当成平衡缺口。</li><li>需求：镀锌、压铸、地产与出口共同验证，避免仅以价格判断终端。</li><li>观点边界：⚙️规则摘要只整理当前数据，不沿用旧周报方向。</li></ul></div></details>')
            parts.append('<div id="sumsec"><h2>总结<span class="en">SUMMARY</span></h2><table class="sum"><tr><th>环节</th><th>相关数据</th><th>周度结论</th><th>跟踪</th></tr>')
            rows = [
                ('价格', '沪锌 %s 元/吨（%s）｜LME锌 %s 美元/吨（%s）' % (shared.fmt(latest['shfe'][1]), latest['shfe'][0], shared.fmt(latest['lme'][1]), latest['lme'][0]), '近5个交易日 %s，价格%s' % (shared.pct(chg5), price_word), price_word, 'nt'),
                ('海外库存', 'LME锌库存 %s 吨（%s）' % (shared.fmt(last(charts['a10134450'])[1]), last(charts['a10134450'])[0]), '近5个交易日 %s，海外%s' % (shared.pct(stock_chg), stock_word), stock_word, 'up' if stock_word == '去库' else 'dn'),
                ('国内库存', '锌锭社库 %s 万吨（%s）' % (shared.fmt(latest['social_stock'][1]), latest['social_stock'][0]), '与LME、SHFE仓单共同验证，不用单一口库存代全部供需', '跟踪', 'nt'),
                ('供应·矿端', '昆明50%% TC %s 元/吨（%s）' % (shared.fmt(latest['tc'][1]), latest['tc'][0]), 'TC、进口矿与精炼锌产量共同判断矿端松紧', '重点', 'nt'),
                ('需求·镀锌', '镀锌板卷产量 %s 万吨（%s）' % (shared.fmt(last(charts['ID00187499'])[1]), last(charts['ID00187499'])[0]), '镀锌、压铸、地产和出口共同验证', '跟踪', 'nt'),
                ('策略观察', '价格、库存、TC、进口盈亏交叉验证', '信号未同向前不作单边外推；优先看库存趋势与现货结构是否持续', '观望', 'nt'),
            ]
            for label, datum, conclusion, status, cls in rows:
                parts.append('<tr><td class="c">%s</td><td>%s</td><td>%s</td><td><span class="dir %s">%s</span></td></tr>' % (
                    shared.esc(label), shared.esc(datum), shared.esc(conclusion), cls, shared.esc(status)))
            parts.append('</table></div><footer>数据更新至 %s ｜ 锌研究周报 ｜ ⚙️自动构建，仅供研究参考。</footer>' % shared.esc(asof))
        elif sid == 'framework':
            parts.append('''<details class="fold" id="framework" open><summary>锌产业研究框架（三条线）</summary><div class="foldbody"><div class="lead">矿端线：锌精矿 → TC/进口 → 冶炼</div><p>跟踪国产与进口TC、港口矿库存、进口矿来源和精炼锌产量，判断矿端议价与冶炼开工。</p><div class="lead">锭端线：精炼锌 → LME/SHFE/社库 → 现货升贴水与月差</div><p>跟踪全球仓单、国内社库、0#锌升贴水、LME 0-3和进口盈亏，区分近月结构与单边方向。</p><div class="lead">需求线：镀锌/压铸/氧化锌 → 地产、汽车、家电与出口</div><p>跟踪镀锌板卷、压铸锌合金、氧化锌开工、地产与出口，避免只用价格反推需求。</p><p><a href="https://amiya866.github.io/metals-framework/">打开金属研究框架 ↗</a> · <a href="https://amiya866.github.io/caibao-hub/">打开锌企财报 ↗</a></p></div></details>''')
        elif sid == 'cost':
            parts.append('<div class="lead">成本与进口盈亏</div>')
            parts.append(shared.chart_group(charts, keys))
            wb = load_workbook(FINANCIAL, read_only=True, data_only=True)
            for sheet_name, limit in [('锌矿企业·季度产量', 12), ('锌锭冶炼企业·季度产量', 10)]:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                headers = list(rows[0])
                period = '26Q3' if '26Q3' in headers else '26Q2'
                idx = headers.index(period)
                guide_idx = headers.index('FY2026指引')
                parts.append('<div class="lead">%s · %s</div><div class="tablewrap"><table class="bal"><tr><th>公司</th><th>国家/项目</th><th>%s</th><th>FY2026指引</th></tr>' % (sheet_name, period, period))
                for row in rows[1:limit + 1]:
                    if row[idx] in (None, '-'):
                        continue
                    parts.append('<tr><td>%s</td><td>%s</td><td class="num">%s</td><td>%s</td></tr>' % (
                        shared.esc(row[0]), shared.esc(row[2] or row[1]), shared.esc(row[idx]), shared.esc(row[guide_idx] or '—')))
                parts.append('</table></div>')
            wb.close()
        elif sid == 'technical':
            parts.append('<div class="tech-grid">')
            cards = [
                ('日线趋势', '偏多' if current_rsi and current_rsi > 55 else ('偏空' if current_rsi and current_rsi < 45 else '震荡'),
                 '收盘 %s；MA5 %s / MA20 %s；RSI14 %.1f；价格%s。' % (shared.fmt(shfe_values[-1]), shared.fmt(ma5), shared.fmt(ma20), current_rsi or 0, price_word),
                 'up' if ma5 > ma20 else 'down'),
                ('MACD动量', '偏多' if macd[-1] > signal[-1] else '偏空',
                 'DIF %s / DEA %s；柱体 %s。' % (shared.fmt(macd[-1]), shared.fmt(signal[-1]), shared.fmt(macd[-1] - signal[-1])),
                 'up' if macd[-1] > signal[-1] else 'down'),
                ('关键区间', '观察', '近20日低 %s / 高 %s；结合库存和升贴水验证，不以单一技术位直接下单。' % (shared.fmt(min(shfe_values[-20:])), shared.fmt(max(shfe_values[-20:]))), 'neutral'),
            ]
            for frame, status, detail, tone in cards:
                parts.append('<article class="tech-card %s"><h3>%s<span class="score">%s</span></h3><p>%s</p></article>' % (tone, frame, status, shared.esc(detail)))
            parts.append('</div><div class="tech-note">技术分析由最新沪锌日线自动计算，仅供研究参考。</div><div class="lead">沪锌主连 · 日线</div>')
            parts.append(shared.chart_group(charts, keys))
        elif sid == 'sources':
            summary = freshness['summary']
            parts.append('<details class="fold" id="datagap" open><summary>数据口径、缺口与新鲜度（%d项全量审计）</summary><div class="foldbody">' % summary['total'])
            parts.append('<div class="fresh-grid">')
            for code, label in [('fresh', '正常'), ('watch', '关注'), ('stale', '滞后'), ('missing', '缺失'), ('review', '待核')]:
                parts.append('<button type="button" class="fresh-kpi %s" data-filter="%s"><span>%s</span><strong>%d</strong><small>点击筛选</small></button>' % (code, code, label, summary[code]))
            parts.append('<button type="button" class="fresh-kpi total" data-filter="all"><span>全部指标</span><strong>%d</strong><small>重置筛选</small></button></div>' % summary['total'])
            parts.append('<p>来源：数据缓存_v1.db、guan/新浪行情缓存、锌企财报真源。日度 7/14 天、周度 10/21 天、月度 45/75 天、季度 150/240 天为正常/关注阈值；缺失和滞后不补零，不以旧值冒充最新。</p></div></details>')
            parts.append('<div class="fresh-legend">当前显示 <b id="freshCount">%d</b> / %d 项｜点击上方卡片筛选</div>' % (summary['total'], summary['total']))
            parts.append('<div class="tablewrap"><table class="bal fresh-table"><tr><th>模块/指标</th><th>来源</th><th>频率</th><th>单位</th><th>最新标签/值</th><th>新鲜度</th><th>处理</th></tr>')
            for row in freshness['rows']:
                age = '，%d天' % row['age_days'] if row['age_days'] is not None else ''
                parts.append('<tr data-fresh="%s"><td><b>%s</b><br><span class="src-note">%s｜%s</span></td><td>%s</td><td>%s</td><td>%s</td><td>%s<br><span class="src-note">%s%s</span></td><td><span class="fresh-badge %s">%s</span></td><td>%s</td></tr>' % (
                    row['status'], shared.esc(row['name']), shared.esc(row['module']), shared.esc(row['key']), shared.esc(row['source']),
                    shared.esc(row['frequency']), shared.esc(row['unit']), shared.esc(row['latest']), shared.esc(row['latest_value']), shared.esc(age),
                    row['status'], shared.esc(row['status_label']), shared.esc(row['action'])))
            parts.append('</table></div>')
        else:
            parts.append('<div class="lead">%s · 最新读数逐项标注日期</div>' % shared.esc(title))
            parts.append(shared.chart_group(charts, keys))
        parts.append('</div></section>')

    payload = json.dumps(charts, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    doc = '''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow,noarchive"><link href="data:," rel="icon">
<title>研究周报 · 锌</title><style>__BASE_CSS__</style><link href="weekly_mobile.css" rel="stylesheet"><link href="weekly_polish.css" rel="stylesheet"><link href="research.css" rel="stylesheet"><link href="../weekly-light.css" rel="stylesheet"><style>.chart-box{height:240px;position:relative}.chart-box canvas{display:block;width:100%;height:100%}.tablewrap{overflow-x:auto}.page-body p{font-size:14.5px}.cbox .stats{line-height:1.55}.fresh-grid{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:9px;margin:12px 0}.fresh-kpi{border:1px solid var(--line);background:var(--cardbg);padding:10px;text-align:left;cursor:pointer}.fresh-kpi span,.fresh-kpi small{display:block;color:var(--faint);font-size:11px}.fresh-kpi strong{display:block;font-size:25px}.fresh-kpi.fresh strong{color:#188865}.fresh-kpi.watch strong{color:#a3760e}.fresh-kpi.stale strong,.fresh-kpi.missing strong{color:#c44b5d}.fresh-kpi.review strong{color:#7c4dff}.fresh-kpi.active{outline:2px solid var(--blue);outline-offset:1px}.fresh-legend{margin:12px 0;color:var(--grey);font-size:13px}.fresh-badge{display:inline-block;padding:1px 7px;border:1px solid var(--line);white-space:nowrap}.fresh-badge.fresh{color:#188865}.fresh-badge.watch{color:#a3760e}.fresh-badge.stale,.fresh-badge.missing{color:#c44b5d}.fresh-badge.review{color:#7c4dff}.src-note{color:var(--faint);font-size:11px}.fresh-table{min-width:1080px}.fresh-table td{vertical-align:top}@media(max-width:900px){.fresh-grid{grid-template-columns:repeat(3,1fr)}}</style></head>
<body><header class="research-topbar"><div class="research-brand"><a aria-label="返回整合研究站" class="back-home" href="../index.html">← 整合研究站</a><div class="brand-copy"><p class="research-eyebrow">COMMODITY WEEKLY RESEARCH · ZINC</p><h1>研究周报 <span class="commodity-current">锌</span></h1><p>矿端TC · 精炼锌 · 全球库存 · 镀锌需求 · 成本平衡</p></div></div><div class="research-actions"><nav class="commodity-switch" aria-label="品种切换"><a href="../tin-weekly/#sumsec">锡</a><a href="../nickel-weekly/#sumsec">镍</a><a class="active" href="#sumsec">锌</a></nav><a class="back-home" href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">锌企财报 ↗</a><span class="research-stamp">数据截至 <b>__ASOF__</b></span><button aria-pressed="false" class="view-toggle" id="viewToggle" type="button">切换手机版</button></div></header>
<div class="research-workspace"><aside class="research-sidebar" id="researchSidebar"><button aria-controls="treeNav" aria-expanded="false" class="menu-toggle" id="menuToggle" type="button">研究目录 <span>＋</span></button><div class="sidebar-heading"><span>RESEARCH INDEX</span><b>研究目录</b></div><nav aria-label="锌研究分层目录" class="tree-nav" id="treeNav"></nav></aside><div class="research-content"><div aria-live="polite" class="breadcrumb" id="breadcrumb"></div><div class="wrap" id="reportWrap">__SECTIONS__</div><section aria-live="polite" class="tool-page" hidden id="toolPage"><div class="tool-head"><div><p class="research-eyebrow">RESEARCH WORKBENCH</p><h2 id="toolTitle">研究工作台</h2></div><button id="closeTool" type="button">← 返回周报</button></div><iframe hidden id="toolFrame" title="锌研究工作台"></iframe></section></div></div>
<script src="lib/chart.min.js"></script><script>const charts=__PAYLOAD__;document.querySelectorAll('canvas[data-chart]').forEach(function(canvas){const c=charts[canvas.dataset.chart];if(!c)return;const ds=(c.datasets||[]).map(function(d){return Object.assign({borderWidth:1.5,pointRadius:0,spanGaps:true},d);});const cfg={type:'line',data:{labels:c.labels||[],datasets:ds},options:{responsive:true,maintainAspectRatio:false,animation:false,interaction:{mode:'index',intersect:false},plugins:{legend:{position:'bottom',labels:{boxWidth:12,font:{family:'KaiTi,"楷体",serif',size:10}}}},scales:{x:{ticks:{maxTicksLimit:8,font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}},y:{ticks:{font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}}}}};new Chart(canvas,cfg);});</script><script src="weekly_mobile.js"></script><script src="research.js"></script><script>document.querySelectorAll('.fresh-kpi').forEach(function(btn){btn.addEventListener('click',function(){const code=btn.dataset.filter;let shown=0;document.querySelectorAll('.fresh-kpi').forEach(function(x){x.classList.toggle('active',x===btn)});document.querySelectorAll('tr[data-fresh]').forEach(function(row){const visible=code==='all'||row.dataset.fresh===code;row.hidden=!visible;if(visible)shown++});document.getElementById('freshCount').textContent=shown})});</script></body></html>'''
    doc = doc.replace('__BASE_CSS__', base_css).replace('__ASOF__', shared.esc(asof)).replace('__SECTIONS__', ''.join(parts)).replace('__PAYLOAD__', payload)
    tmp = OUT / 'index.html.tmp'
    tmp.write_text(doc, encoding='utf-8')
    os.replace(tmp, OUT / 'index.html')
    summary = {'title': '研究周报 · 锌', 'asof': asof, 'headline': headline, 'metrics': metrics,
               'local_url': 'zinc-weekly/#sumsec', 'source_url': 'https://amiya866.github.io/caibao-hub/',
               'charts': len(charts), 'freshness': freshness['summary'], 'sync_status': 'ok'}
    (BASE / 'data' / 'zinc_weekly.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    update_daily(summary)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == '__main__':
    main()

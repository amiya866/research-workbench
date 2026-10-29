# -*- coding: utf-8 -*-
"""由本地缓存、行情和铝企财报生成与锡镍同结构的铝研究周报。"""
import csv
import datetime
import json
import os
import re
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

import build_nickel_weekly as shared

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'aluminum-weekly'
CACHE = Path(r'D:/Kimi/金属总网/网站构建/db/数据缓存_v1.db')
SHFE_CSV = Path(r'D:/Kimi/有色截面配置_2026/_价格缓存.csv')
LME_CSV = Path(r'D:/Kimi/有色截面配置_2026/_LME价格缓存.csv')
FINANCIAL = Path(r'D:/Kimi/金属总网/网站构建/财报汇总/铝/全球铝企季度产量梳理.xlsx')
META = BASE / 'data' / 'aluminum_weekly_freshness.json'

SECTIONS = [
    ('sumsec', '本周总结', '总览', []),
    ('price', '价格与结构', '总览', ['overview_shfe', 'overview_lme', 'overview_lme03', 'FU00024308', 'ID00188139']),
    ('supply_mine', '铝土矿与氧化铝', '重点关注/供给', ['a10016324', 'a10016284', 'a10016292', 'a10016300', 'a10015944', 'a10167792', 'ID01001945', 'a10031611']),
    ('supply_smelt', '电解铝', '重点关注/供给', ['a10124317', 'ID01001838', 'ID01001837', 'a10016193', 'ID01001759', 'ID01131701']),
    ('inventory', '显性库存', '重点关注/供给', ['overview_global_stock', 'ID00188281', 'a10134438', 'a10166772', 'ID00188307', 'ID00188313']),
    ('demand', '铝型材/板带/终端', '重点关注/需求', ['a10031808', 'a10016011', 'a12750799', 'a10102727', 'a10102770', 'ID00275563']),
    ('cost', '成本与利润', '平衡表', ['a10016216', 'a10016215', 'a10016211', 'a10016206', 'ID00188177', 'ID01001761']),
    ('balance', '年度平衡与2027策略', '平衡表', []),
    ('framework', '产业框架', '总览', []),
    ('technical', '技术分析', '宏观与策略', ['overview_shfe']),
    ('sources', '数据缺口', '总览', []),
]
NAMES = {key: title for key, title in []}
UNITS = {'overview_shfe': '元/吨', 'overview_lme': '美元/吨', 'overview_global_stock': '吨'}
KPI = [('shfe', '沪铝主连', '元/吨'), ('lme', 'LME铝', '美元/吨'), ('lme03', 'LME 0-3', '美元/吨'),
       ('global_stock', 'LME+SHFE库存', '吨'), ('social_stock', '铝锭社库', '万吨'), ('tc', '电解铝利润', '元/吨'),
       ('domestic', 'SHFE铝基差', '元/吨'), ('alumina', '氧化铝价格', '元/吨')]


def series(conn, ind_id, name, unit):
    rows = conn.execute('select date,value from series_points where ind_id=? order by date', (ind_id,)).fetchall()
    return {'labels': [r[0] for r in rows], 'datasets': [{'label': name, 'data': [r[1] for r in rows],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}, unit


def csv_chart(path, column, name):
    labels, values = [], []
    with open(path, encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            value = row.get(column)
            if value not in (None, ''):
                labels.append(row[''])
                values.append(float(value))
    return {'labels': labels, 'datasets': [{'label': name, 'data': values, 'borderColor': '#1F4E79',
            'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def cached_lme03_chart():
    path = BASE / 'data' / 'lme_cash3m.json'
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding='utf-8')).get('AL', {})
    hist = data.get('hist') or {}
    rows = sorted((date, values[0] - values[1]) for date, values in hist.items() if isinstance(values, list) and len(values) >= 2)
    if not rows:
        return None
    return {'labels': [row[0] for row in rows], 'datasets': [{'label': 'LME铝 0-3 升贴水', 'data': [row[1] for row in rows],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def last(chart):
    point = shared.last_point(chart)
    return [point[0], point[1]] if point else [None, None]


def empty_chart():
    return {'labels': [], 'datasets': [], '_type': 'line'}


def recover_previous_charts(charts, source_meta):
    warnings = []
    path = OUT / 'index.html'
    if not path.exists():
        return warnings
    match = re.search(r'<script>const charts=(.*?);document\.querySelectorAll', path.read_text(encoding='utf-8'), re.S)
    if not match:
        return warnings
    try:
        previous = json.loads(match.group(1))
    except json.JSONDecodeError:
        return warnings
    for key, old_chart in previous.items():
        old_point = shared.last_point(old_chart)
        if old_point and not shared.last_point(charts.get(key, {})):
            charts[key] = old_chart
            source_meta[key] = dict(source_meta.get(key) or {}, recoveredFrom=str(old_point[0]))
            warnings.append('%s 当前源无有效数据，保留上一成功版（%s）' % (key, old_point[0]))
    return warnings


def global_stock(lme, shfe):
    if not shared.last_point(lme) or not shared.last_point(shfe):
        return empty_chart()
    dates = sorted(set(lme['labels']) | set(shfe['labels']))
    lv, sv = lme['datasets'][0]['data'], shfe['datasets'][0]['data']
    lm = dict(zip(lme['labels'], lv))
    sm = dict(zip(shfe['labels'], sv))
    values, last_l, last_s = [], None, None
    for date in dates:
        last_l = lm.get(date, last_l)
        last_s = sm.get(date, last_s)
        values.append(last_l + last_s if last_l is not None and last_s is not None else None)
    return {'labels': dates, 'datasets': [{'label': 'LME+SHFE库存', 'data': values,
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def rsi(values, period=14):
    if len(values) <= period:
        return None
    changes = [b - a for a, b in zip(values[-period - 1:-1], values[-period:])]
    gains = sum(max(x, 0) for x in changes) / period
    losses = sum(max(-x, 0) for x in changes) / period
    return 100 if not losses else 100 - 100 / (1 + gains / losses)


def ema(values, period):
    alpha = 2 / (period + 1)
    out = []
    for i, value in enumerate(values):
        out.append(value if i == 0 else alpha * value + (1 - alpha) * out[-1])
    return out


def update_daily(summary):
    path = BASE / 'reports' / 'daily_latest.html'
    if not path.exists():
        return
    page = path.read_text(encoding='utf-8')
    begin, end = '<!-- ALUMINUM_WEEKLY_CARD_BEGIN -->', '<!-- ALUMINUM_WEEKLY_CARD_END -->'
    metrics = ''.join('<li><b>%s</b>　%s</li>' % (shared.esc(m['name']), shared.esc(m['value'])) for m in summary['metrics'][:4])
    card = '''%s
<div class="card" id="aluminumWeeklyCard" style="border-left:3px solid var(--gold);margin:14px 0;">
  <div style="display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap;align-items:center;">
    <h3 style="margin:0;color:var(--gold);">研究周报 · 铝</h3>
    <a href="../aluminum-weekly/#sumsec" style="font-weight:700;">打开完整铝周报 ↗</a>
  </div>
  <p style="margin-top:8px;">%s</p>
  <ul style="margin:8px 0 0 18px;font-size:12px;color:var(--dim);">%s</ul>
  <p style="font-size:11px;color:var(--dim);margin-top:8px;">数据截至 %s · 本地生成，随研究站管线自动更新 · <a href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">铝企财报</a></p>
</div>
%s''' % (begin, shared.esc(summary['headline']), metrics, shared.esc(summary['asof']), end)
    pattern = re.escape(begin) + r'.*?' + re.escape(end)
    if re.search(pattern, page, re.S):
        page = re.sub(pattern, lambda _m: card, page, count=1, flags=re.S)
    else:
        anchor = '<!-- NICKEL_WEEKLY_CARD_END -->'
        if anchor not in page:
            anchor = '<section class="dp on" id="dp1">'
        page = page.replace(anchor, anchor + '\n' + card, 1)
    path.write_text(page, encoding='utf-8')


def annual_balance_html(shfe_values, full=False):
    window = shfe_values[-120:]
    low, high, current = min(window), max(window), shfe_values[-1]
    position = (current - low) / (high - low) if high > low else 0.5
    rows = [
        ('2025实际', '7,405.6', '7,420.1', '-14.4', '小幅短缺'),
        ('2026E', '7,442.7（+0.5%）', '7,530.7（+1.5%）', '-88.0', '短缺扩大，近月不能空'),
        ('2027E', '7,810.6（+4.9%）', '7,757.0（+3.0%）', '+53.5', '转为过剩，远月空配窗口'),
    ]
    table = ''.join('<tr><td>%s</td><td class="num">%s</td><td class="num">%s</td><td class="num">%s</td><td>%s</td></tr>' % row for row in rows)
    parts = ['<details class="fold" id="annualbalance%s" open><summary>年度平衡与2027空配：2026短缺 → 2027过剩</summary><div class="foldbody">' % ('Full' if full else ''),
             '<div class="lead">全球原铝供需平衡（百万吨→万吨展示）</div>',
             '<div class="tablewrap"><table class="bal"><tr><th>年份</th><th>产量（万吨）</th><th>消费（万吨）</th><th>平衡（万吨）</th><th>交易含义</th></tr>', table, '</table></div>',
             '<p><b>区域拆分：</b>中国平衡2026E +22.4万吨→2027E -6.7万吨；海外（除中国）平衡2026E -110.4万吨→2027E <b>+60.2万吨</b>。2027过剩主要由海外复产与新增供应驱动，不能只看国内产能天花板。</p>',
             '<p class="src-note">来源：CRU Aluminium Market Outlook 2026-06（2025实际、2026/2027预测）；2026Q3公司产量、库存与复产信息用于滚动验证。预测不是实际值，年度平衡需随新数据重拍。</p>',
             '<div class="lead">可执行策略：空2027远月，不空2026近月</div><ul class="sq">',
             '<li><b>方向：</b>2026短缺约束近月下跌空间；2027转过剩是中期空配依据。优先2027年远月反弹空，近月低库存/Back时不做长期裸空。</li>',
             '<li><b>入场：</b>远月反弹至各自120日区间80%%分位以上，且现货升水不再走强、氧化铝/海外复产继续兑现，分批建立空配。当前沪铝主连120日位置为 %s（%s，区间%s—%s），主连位置仅作温度参考，执行须看对应远月合约。</li>' % (shared.pct(position), shared.fmt(current), shared.fmt(low), shared.fmt(high)),
             '<li><b>风控：</b>单笔远月止损参考入场价上方3%—4%或突破对应合约120日高点；若海外复产再中断、库存连续两个月加速下降或LME Back持续走强，停止加空并减仓。</li>',
             '<li><b>止盈/复核：</b>第一目标为远月回落5%—8%，第二目标10%—15%，到达后分批兑现而非无限期持有；每月重拍年度平衡。若2027平衡被下修至短缺，长期空配逻辑失效。</li>',
             '<li><b>更稳结构：</b>近月强Back时可研究“多近月、空2027远月”的时间结构，但两腿合约流动性和保证金不同，不当作无风险套利。</li></ul>']
    if full:
        parts += ['<div class="lead">2027平衡敏感性（研究情景，不是CRU原预测）</div>',
                  '<div class="tablewrap"><table class="bal"><tr><th>情景</th><th>产量增速</th><th>消费增速</th><th>2027平衡</th><th>策略</th></tr>',
                  '<tr><td>紧张</td><td class="num">+3.5%</td><td class="num">+4.0%</td><td class="num">-128.7万吨</td><td>停止空配</td></tr>',
                  '<tr><td>基准</td><td class="num">+4.9%</td><td class="num">+3.0%</td><td class="num">+53.5万吨</td><td>远月逢高空</td></tr>',
                  '<tr><td>宽松</td><td class="num">+5.5%</td><td class="num">+2.0%</td><td class="num">+170.7万吨</td><td>提高空配优先级</td></tr></table></div>',
                  '<div class="lead">供应增量路径与需要盯的证伪点</div><ul class="sq">',
                  '<li><b>中东复产：</b>Alba受损产线修复、EGA恢复是海外增量的重要组成；若再次中断，2027过剩下修。</li>',
                  '<li><b>欧洲/北美：</b>San Ciprián复产、Arvida旧线退出与AP60爬产构成置换；净增量不能只看名义产能。</li>',
                  '<li><b>印度与氧化铝：</b>Vedanta电解铝和Lanjigarh氧化铝创纪录爬产，若氧化铝价格继续下降，电解铝高利润会进一步刺激供应。</li>',
                  '<li><b>反向约束：</b>Mozal停产、Bell Bay永久退出抵消部分增量；2027是否过剩要看复产净额，而不是把新增项目简单相加。</li></ul>']
    parts.append('</div></details>')
    return ''.join(parts)


def main():
    conn = sqlite3.connect(f'file:{CACHE.as_posix()}?mode=ro', uri=True)
    charts, source_meta, names, units = {}, {}, dict(NAMES), dict(UNITS)
    for ind_id, name, unit in conn.execute('select ind_id,name,unit from series_meta where comm="AL"'):
        charts[ind_id], units[ind_id] = series(conn, ind_id, name, unit)
        source_meta[ind_id] = {'type': 'zhiji', 'indicator': ind_id, 'unit': unit}
        names[ind_id] = name
    conn.close()
    charts['overview_shfe'] = csv_chart(SHFE_CSV, 'AL', '沪铝主连')
    charts['overview_lme'] = csv_chart(LME_CSV, 'AL', 'LME铝')
    basis_chart = cached_lme03_chart()
    charts['overview_lme03'] = basis_chart or empty_chart()
    units['overview_lme03'] = '美元/吨'
    source_meta['overview_lme03'] = {'type': 'derived', 'formula': 'westmetall cash-3M'}
    names['overview_lme03'] = 'LME铝 0-3 升贴水'
    charts['overview_global_stock'] = global_stock(charts['a10134438'], charts['ID00188281'])
    for key, label in [('overview_shfe', 'guan+本地日K'), ('overview_lme', '新浪外盘+本地缓存'), ('overview_global_stock', 'LME+SHFE合计')]:
        source_meta[key] = {'type': 'derived', 'formula': label}
        names[key] = {'overview_shfe': '沪铝主连', 'overview_lme': 'LME铝', 'overview_global_stock': '全球显性库存'}[key]
    if not shared.last_point(charts['overview_shfe']):
        raise RuntimeError('核心沪铝行情无有效数据，保留上一成功版并中止重建')
    warnings = recover_previous_charts(charts, source_meta)
    if not shared.last_point(charts['overview_shfe']):
        raise RuntimeError('核心沪铝行情当前与上一成功版均无有效数据，中止重建')
    raw_charts = charts
    referenced = {key for _sid, _title, _trail, keys in SECTIONS for key in keys}
    empty = sorted(key for key in referenced if not shared.last_point(charts.get(key, {})))
    shared.NAMES = names
    shared.SECTIONS = SECTIONS
    built = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
    data = {'charts': charts, 'sourceMeta': {'charts': source_meta}, 'meta': {'builtAt': built, 'warnings': warnings, 'emptyCharts': empty}}
    freshness = shared.build_freshness(data)
    META.write_text(json.dumps(freshness, ensure_ascii=False, indent=2), encoding='utf-8')
    latest = {
        'shfe': last(raw_charts['overview_shfe']), 'lme': last(raw_charts['overview_lme']),
        'lme03': last(raw_charts['overview_lme03']), 'domestic': last(raw_charts['FU00024308']),
        'global_stock': last(raw_charts['overview_global_stock']),
        'social_stock': last(raw_charts['a10166772']), 'tc': last(raw_charts['a10016216']),
        'alumina': last(raw_charts['ID00188139']),
    }
    asof = max(value[0] for value in latest.values() if value[0])
    shfe_values = [v for v in charts['overview_shfe']['datasets'][0]['data'] if v is not None]
    chg5 = shfe_values[-1] / shfe_values[-6] - 1 if len(shfe_values) > 6 else None
    stock_values = [v for v in charts['a10134438']['datasets'][0]['data'] if v is not None]
    stock_chg = stock_values[-1] / stock_values[-6] - 1 if len(stock_values) > 6 else None
    ma5 = sum(shfe_values[-5:]) / 5
    ma20 = sum(shfe_values[-20:]) / 20
    current_rsi = rsi(shfe_values)
    macd = [a - b for a, b in zip(ema(shfe_values, 12), ema(shfe_values, 26))]
    signal = ema(macd, 9)
    price_word = '偏强' if (chg5 or 0) > 0.01 else ('偏弱' if (chg5 or 0) < -0.01 else '震荡')
    stock_word = '累库' if (stock_chg or 0) > 0 else '去库'
    lme_structure = 'Back' if (latest['lme03'][1] or 0) > 0 else ('Contango' if (latest['lme03'][1] or 0) < 0 else '持平')
    domestic_structure = '升水' if (latest['domestic'][1] or 0) > 0 else ('贴水' if (latest['domestic'][1] or 0) < 0 else '持平')
    structure_desc = 'LME 0-3 %s 美元/吨（%s），SHFE铝基差%s %s 元/吨（%s）' % (
        shared.fmt(latest['lme03'][1]), latest['lme03'][0], domestic_structure,
        shared.fmt(latest['domestic'][1]), latest['domestic'][0])
    headline = ('本周定调：沪铝%s与LME%s并存，结构%s —— 沪铝近5个交易日 %s，LME库存近5个交易日 %s；'
                '铝锭社库 %s 万吨（%s），氧化铝价格 %s 元/吨（%s）；结构：%s。'
                '近端重点跟踪价格、结构、库存、氧化铝与电解铝利润能否同向验证；⚙️规则摘要，不构成投资建议。') % (
        price_word, stock_word, lme_structure, shared.pct(chg5), shared.pct(stock_chg), shared.fmt(latest['social_stock'][1]), latest['social_stock'][0],
        shared.fmt(latest['alumina'][1]), latest['alumina'][0], structure_desc)
    metrics = [{'name': name, 'value': '%s %s（%s）' % (shared.fmt(latest[key][1]), unit, latest[key][0])} for key, name, unit in KPI]
    q3 = shared.period_return(charts['overview_shfe'], '2026-07-01', '2026-09-30')
    q4 = shared.period_return(charts['overview_shfe'], '2026-10-01', asof)
    output = shared.last_point(charts['a10124317'])
    downstream = shared.last_point(charts['a10031808'])
    q4_html = '''<details class="fold" id="q4outlook"><summary>季度展望：Q3复盘与Q4展望（点击展开/折叠）</summary><div class="foldbody">
<div class="lead">一、三季度复盘与四季度展望</div><ul class="sq">
<li>Q3沪铝价格变动 %s，氧化铝与电解铝供应恢复、利润高企和库存偏低共同塑造价格弹性，成本高利润与弱宏观之间的分歧加大。</li>
<li>Q4迄今价格变动 %s。当前LME与SHFE合计库存 %s 吨、国内社库 %s 万吨，低库存仍支撑近月，但氧化铝价格 %s 元/吨与电解铝利润 %s 元/吨说明供应端利润已较充裕。</li>
<li>结构定位：LME 0-3 %s 美元/吨（%s），SHFE铝基差 %s 元/吨（%s）；期现结构必须与复产和库存方向一致才有效。</li>
<li>基准判断是“低库存支撑近月，高利润刺激供应恢复”。Q4重点不是单看铝价强弱，而是核查氧化铝、电解铝复产能否改变库存和现货结构。</li></ul>
<div class="lead">二、供应：三层产量与利润恢复</div><ul class="sq">
<li>铝土矿、氧化铝和电解铝不能混为一层：进口矿和氧化铝恢复决定成本端，运行产能和铝水比例决定锭端实际供应。</li>
<li>中国电解铝最新 %s 万吨（%s）；美铝与Vedanta披露的26Q3原铝产量均为64.9万吨，海外复产也在贡献增量。</li>
<li>高利润有利于复产和爬产，但若氧化铝价格下跌快于铝价，成本支撑会同步弱化，需防止利润高位后的供应回归。</li></ul>
<div class="lead">三、需求与库存：铝加工托底，出口验证</div><ul class="sq">
<li>铝型材龙头开工率最新 %s%%（%s），板带、地产竣工、家电和铝材出口共同决定真实订单，不宜只用价格判断消费。</li>
<li>若社库继续下降且下游开工稳定，近月结构仍有支撑；若复产叠加出口走弱，库存拐点将首先削弱远月估值。</li>
<li>氧化铝与电解铝利润需要分开跟踪，避免把上游价格回落简单理解为全行业利空。</li></ul>
<div class="lead">四、策略与验证点</div><ul class="sq">
<li>近月看库存、基差和LME 0-3，远月看复产与利润回归，避免在高利润阶段仅凭低库存追涨。</li>
<li>上修条件：持续去库＋结构走强＋铝加工开工稳定＋氧化铝供应扰动；下修条件：复产超预期＋结构走弱＋出口转弱＋社库回升。</li>
<li>本段为数据驱动的研究框架，不构成投资建议。</li></ul></div></details>''' % (
        shared.pct(q3), shared.pct(q4), shared.fmt(latest['global_stock'][1]), shared.fmt(latest['social_stock'][1]),
        shared.fmt(latest['alumina'][1]), shared.fmt(latest['tc'][1]),
        shared.fmt(latest['lme03'][1]), latest['lme03'][0], shared.fmt(latest['domestic'][1]), latest['domestic'][0],
        shared.fmt(output[1]) if output else '—', output[0] if output else '—',
        shared.fmt(downstream[1]) if downstream else '—', downstream[0] if downstream else '—')
    q4_html = shared.quarterly_html(OUT / 'index.html', q4_html)

    def _sign(x):
        return 1 if (x or 0) > 0 else (-1 if (x or 0) < 0 else 0)
    global20 = shared.change(charts['overview_global_stock'], 20)
    social20 = shared.change(charts['a10166772'], 20)
    output_chg = shared.change(charts['a10124317'], 20)
    downstream_chg = shared.change(charts['a10031808'], 20)
    lme03_trend = shared.change(charts['overview_lme03'], 20)
    domestic_trend = shared.change(charts['FU00024308'], 20)
    lme_s = _sign(latest['lme03'][1]); dom_s = _sign(latest['domestic'][1])
    if lme_s and dom_s and lme_s == dom_s:
        structure_score = lme_s * 2; structure_judge = '内外结构同向支撑（方向较明确）' if structure_score > 0 else '内外结构同向偏弱（方向较明确）'
    elif lme_s or dom_s:
        structure_score = lme_s or dom_s; structure_judge = '单边结构成立，仍需另一边确认'
    else:
        structure_score = 0; structure_judge = '结构数据不足，跟踪'
    if global20 is not None and social20 is not None and _sign(global20) == _sign(social20):
        inventory_score = -_sign(global20); inventory_judge = '内外库存同向去化，近月有支撑' if inventory_score > 0 else '内外库存同向累积，近月承压'
    else:
        inventory_score = 0; inventory_judge = '库存信号分化，跟踪'
    supply_score = -_sign(output_chg)
    supply_judge = '电解铝产量回落，供应端有支撑' if supply_score > 0 else ('电解铝产量恢复，供应端施压' if supply_score < 0 else '产量变化不显著，跟踪复产')
    demand_score = 1 if (downstream_chg or 0) > 0.02 else (-1 if (downstream_chg or 0) < -0.02 else 0)
    demand_judge = '铝加工开工改善，偏多证据' if demand_score > 0 else ('铝加工开工走弱，偏空证据' if demand_score < 0 else '需求变化不足，跟踪')
    marginal_items = [
        ('价格', '近5个交易日 %s' % shared.pct(chg5), '价格上行，偏多证据' if (chg5 or 0) > 0 else ('价格下行，偏空证据' if (chg5 or 0) < 0 else '价格震荡，等待方向'), _sign(chg5)),
        ('结构', structure_desc, structure_judge, structure_score),
        ('库存', '全球显性库存20日 %s，社库20日 %s' % (shared.pct(global20), shared.pct(social20)), inventory_judge, inventory_score),
        ('供应', '电解铝产量近20日 %s，电解铝利润 %s 元/吨' % (shared.pct(output_chg), shared.fmt(latest['tc'][1])), supply_judge, supply_score),
        ('需求', '铝型材龙头开工率近20日 %s' % shared.pct(downstream_chg), demand_judge, demand_score),
    ]
    marginal_html = shared.marginal_block(marginal_items)
    weekly_html = shared.weekly_outlook_html(
        '沪铝近5个交易日%s，%s；全球显性库存20日%s、国内社库20日%s。近端低库存和现货结构仍可支撑价格，但供应恢复与加工开工分化限制追涨。' % (shared.pct(chg5), structure_desc, shared.pct(global20), shared.pct(social20)),
        '2027年基准判断是供应增长：国内电解铝受产能天花板约束，增幅更多看运行率与铝水转化；海外新增/复产及铝土矿、氧化铝投放是更重要增量。因而近月紧与远月供应回升可以同时成立，不能把当前低库存线性外推。',
        '近端尊重低库存与升水，不逆势追空；远端不追高，重点等待复产、氧化铝供应和库存拐点确认后布局偏空或近强远弱结构。若持续去库且下游开工改善，则继续保留近月多头思路。',
        '上修：持续去库＋升水/Back走强＋铝加工开工改善。下修：国内外供应增量兑现＋社库连续回升＋氧化铝下跌带动成本下移＋出口或终端订单走弱。')
    charts = shared.seasonalize_charts(charts)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'lib').mkdir(exist_ok=True)
    for name in ('research.css', 'weekly_mobile.css', 'weekly_polish.css', 'weekly_mobile.js'):
        shutil.copyfile(BASE / 'nickel-weekly' / name, OUT / name)
    shutil.copyfile(BASE / 'nickel-weekly' / 'lib' / 'chart.min.js', OUT / 'lib' / 'chart.min.js')
    shutil.copyfile(BASE / 'scripts' / 'aluminum_weekly_research.js', OUT / 'research.js')
    base_css = re.search(r'<style>(.*?)</style>', (BASE / 'nickel-weekly' / 'index.html').read_text(encoding='utf-8'), re.S).group(1)

    parts = []
    for sid, title, trail, keys in SECTIONS:
        parts.append('<section class="report-page" data-page="%s" data-trail="%s"><div class="page-heading"><h2>%s</h2><span>%s</span></div><div class="page-body">' % (
            sid, shared.esc(trail), shared.esc(title), shared.esc(trail.replace('/', ' / ').upper())))
        if sid in ('sumsec', 'balance'):
            parts.append(annual_balance_html(shfe_values, full=(sid == 'balance')))
        if sid == 'sumsec':
            parts.append('<div class="tone"><b>%s</b></div><div class="roll">策略每日滚动更新：周报为周五存档，最新决策见日报页。使用左侧目录切换研究内容。</div>' % shared.esc(headline))
            parts.append(weekly_html)
            parts.append(q4_html)
            parts.append('<div id="sumsec"><h2>总结<span class="en">SUMMARY</span></h2><table class="sum"><tr><th>环节</th><th>相关数据</th><th>周度结论</th><th>跟踪</th></tr>')
            def _cls(score):
                return 'up' if (score or 0) > 0 else ('down' if (score or 0) < 0 else 'nt')
            score_map = {item[0]: item[3] for item in marginal_items}
            if (latest['lme03'][1] or 0) > 0:
                lme_status, lme_cls = 'Back偏多', 'up'
                lme_conclusion = 'Back结构支撑近月' if (lme03_trend or 0) >= 0 else 'Back仍在但20日走弱，跟踪是否转弱'
            elif (latest['lme03'][1] or 0) < 0:
                lme_status, lme_cls = 'Contango偏空', 'down'
                lme_conclusion = 'Contango扩大，海外结构走弱' if (lme03_trend or 0) < 0 else 'Contango收窄，海外结构边际改善，仍需确认'
            else:
                lme_status, lme_cls, lme_conclusion = '数据缺失', 'nt', '海外结构暂不可得，继续跟踪'
            if (latest['domestic'][1] or 0) > 0:
                domestic_status, domestic_cls = ('升水走弱，关注', 'nt') if (domestic_trend or 0) < 0 else ('升水偏多', 'up')
                domestic_conclusion = '国内基差升水但20日走弱，跟踪是否转贴水' if (domestic_trend or 0) < 0 else '国内基差升水，结构仍有支撑'
            elif (latest['domestic'][1] or 0) < 0:
                domestic_status, domestic_cls = ('贴水扩大，偏空', 'down') if (domestic_trend or 0) < 0 else ('贴水收窄，关注', 'nt')
                domestic_conclusion = '国内基差贴水扩大，结构走弱' if (domestic_trend or 0) < 0 else '国内基差贴水收窄，结构边际改善，仍需确认'
            else:
                domestic_status, domestic_cls, domestic_conclusion = '数据缺失', 'nt', '国内基差暂不可得，继续跟踪'
            if social20 is not None:
                social_status = '去库偏多' if social20 < 0 else '累库偏空'; social_cls = 'up' if social20 < 0 else 'down'
                social_conclusion = '国内社库20日 %s，%s' % (shared.pct(social20), '去库支撑近月' if social20 < 0 else '累库压制近月')
            else:
                social_status, social_cls, social_conclusion = '跟踪', 'nt', '社库趋势暂不可得，继续跟踪'
            rows = [
                ('价格', '沪铝 %s 元/吨（%s）｜LME铝 %s 美元/吨（%s）' % (shared.fmt(latest['shfe'][1]), latest['shfe'][0], shared.fmt(latest['lme'][1]), latest['lme'][0]), '近5个交易日 %s，价格%s' % (shared.pct(chg5), price_word), price_word, _cls(chg5)),
                ('结构·LME', 'LME 0-3 %s 美元/吨（%s，20日 %s）' % (shared.fmt(latest['lme03'][1]), latest['lme03'][0], shared.pct(lme03_trend)), lme_conclusion, lme_status, lme_cls),
                ('结构·国内', 'SHFE铝基差 %s 元/吨（%s，20日 %s）' % (shared.fmt(latest['domestic'][1]), latest['domestic'][0], shared.pct(domestic_trend)), domestic_conclusion, domestic_status, domestic_cls),
                ('海外库存', 'LME铝库存 %s 吨（%s）' % (shared.fmt(last(raw_charts['a10134438'])[1]), last(raw_charts['a10134438'])[0]), '近5个交易日 %s，海外%s' % (shared.pct(stock_chg), stock_word), '%s%s' % (stock_word, '偏多' if stock_word == '去库' else '偏空'), 'up' if stock_word == '去库' else 'dn'),
                ('国内库存', '铝锭社库 %s 万吨（%s）' % (shared.fmt(latest['social_stock'][1]), latest['social_stock'][0]), social_conclusion, social_status, social_cls),
                ('供应·电解铝', '电解铝产量近20日 %s｜氧化铝价格 %s 元/吨（%s）' % (shared.pct(output_chg), shared.fmt(last(raw_charts['ID00188139'])[1]), last(raw_charts['ID00188139'])[0]), '氧化铝是成本层，产量恢复才是锭端供应压力', '偏多' if score_map['供应'] > 0 else ('偏空' if score_map['供应'] < 0 else '跟踪'), _cls(score_map['供应'])),
                ('需求·铝加工', '铝型材龙头开工率 %s%%（%s，近20日 %s）' % (shared.fmt(last(raw_charts['a10031808'])[1]), last(raw_charts['a10031808'])[0], shared.pct(downstream_chg)), '铝型材、板带、地产、家电和出口共同验证', '偏多' if score_map['需求'] > 0 else ('偏空' if score_map['需求'] < 0 else '跟踪'), _cls(score_map['需求'])),
                ('策略观察', '价格、LME结构、国内结构、库存、产量和铝加工开工交叉验证', '信号未同向前不作单边外推；国内基差走弱时下调近月判断', '跟踪', 'nt'),
            ]
            for label, datum, conclusion, status, cls in rows:
                parts.append('<tr><td class="c">%s</td><td>%s</td><td>%s</td><td><span class="dir %s">%s</span></td></tr>' % (
                    shared.esc(label), shared.esc(datum), shared.esc(conclusion), cls, shared.esc(status)))
            parts.append('</table></div>')
            parts.append(marginal_html)
            parts.append('<footer>数据更新至 %s ｜ 铝研究周报 ｜ ⚙️自动构建，仅供研究参考。</footer>' % shared.esc(asof))
        elif sid == 'framework':
            parts.append('''<details class="fold" id="framework" open><summary>铝产业研究框架（三条线）</summary><div class="foldbody"><div class="lead">矿端线：铝土矿 → 氧化铝 → 电解铝</div><p>跟踪几内亚等进口矿、氧化铝产量与运行产能、电解铝产量和产能利用率，区分三层供应。</p><div class="lead">锭端线：电解铝 → LME/SHFE/社库 → 基差与月差</div><p>跟踪全球仓单、国内社库、SHFE基差与进出口，区分近月结构与单边方向。</p><div class="lead">需求线：铝型材/板带 → 地产、家电、汽车与出口</div><p>跟踪铝型材和板带开工、地产竣工、空调和铝材出口，避免只用价格反推需求。</p><p><a href="https://amiya866.github.io/metals-framework/">打开金属研究框架 ↗</a> · <a href="https://amiya866.github.io/caibao-hub/">打开铝企财报 ↗</a></p></div></details>''')
        elif sid == 'cost':
            parts.append('<div class="lead">成本与进口盈亏</div>')
            parts.append(shared.chart_group(charts, keys))
            wb = load_workbook(FINANCIAL, read_only=True, data_only=True)
            for sheet_name, limit in [('铝土矿·季度产量', 12), ('电解铝·季度产量', 10)]:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                headers = list(rows[0])
                period = '26Q3' if '26Q3' in headers else '26Q2'
                idx = headers.index(period)
                guide_idx = headers.index('FY2026指引')
                parts.append('<div class="lead">%s · %s</div><div class="tablewrap"><table class="bal"><tr><th>公司</th><th>国家/项目</th><th>%s</th><th>FY2026指引</th></tr>' % (sheet_name, period, period))
                for row in rows[1:limit + 1]:
                    if row[idx] in (None, '-'):
                        continue
                    parts.append('<tr><td>%s</td><td>%s</td><td class="num">%s</td><td>%s</td></tr>' % (
                        shared.esc(row[0]), shared.esc(row[2] or row[1]), shared.esc(row[idx]), shared.esc(row[guide_idx] or '—')))
                parts.append('</table></div>')
            wb.close()
        elif sid == 'technical':
            parts.append('<div class="tech-grid">')
            cards = [
                ('日线趋势', '偏多' if current_rsi and current_rsi > 55 else ('偏空' if current_rsi and current_rsi < 45 else '震荡'),
                 '收盘 %s；MA5 %s / MA20 %s；RSI14 %.1f；价格%s。' % (shared.fmt(shfe_values[-1]), shared.fmt(ma5), shared.fmt(ma20), current_rsi or 0, price_word),
                 'up' if ma5 > ma20 else 'down'),
                ('MACD动量', '偏多' if macd[-1] > signal[-1] else '偏空',
                 'DIF %s / DEA %s；柱体 %s。' % (shared.fmt(macd[-1]), shared.fmt(signal[-1]), shared.fmt(macd[-1] - signal[-1])),
                 'up' if macd[-1] > signal[-1] else 'down'),
                ('关键区间', '观察', '近20日低 %s / 高 %s；结合库存和升贴水验证，不以单一技术位直接下单。' % (shared.fmt(min(shfe_values[-20:])), shared.fmt(max(shfe_values[-20:]))), 'neutral'),
            ]
            for frame, status, detail, tone in cards:
                parts.append('<article class="tech-card %s"><h3>%s<span class="score">%s</span></h3><p>%s</p></article>' % (tone, frame, status, shared.esc(detail)))
            parts.append('</div><div class="tech-note">技术分析由最新沪铝日线自动计算，仅供研究参考。</div><div class="lead">沪铝主连 · 日线</div>')
            parts.append(shared.chart_group(charts, keys))
        elif sid == 'sources':
            summary = freshness['summary']
            parts.append('<details class="fold" id="datagap" open><summary>数据口径、缺口与新鲜度（%d项全量审计）</summary><div class="foldbody">' % summary['total'])
            parts.append('<div class="fresh-grid">')
            for code, label in [('fresh', '正常'), ('watch', '关注'), ('stale', '滞后'), ('missing', '缺失'), ('review', '待核')]:
                parts.append('<button type="button" class="fresh-kpi %s" data-filter="%s"><span>%s</span><strong>%d</strong><small>点击筛选</small></button>' % (code, code, label, summary[code]))
            parts.append('<button type="button" class="fresh-kpi total" data-filter="all"><span>全部指标</span><strong>%d</strong><small>重置筛选</small></button></div>' % summary['total'])
            parts.append('<p>来源：数据缓存_v1.db、guan/新浪行情缓存、铝企财报真源。日度 7/14 天、周度 10/21 天、月度 45/75 天、季度 150/240 天为正常/关注阈值；缺失和滞后不补零，不以旧值冒充最新。</p></div></details>')
            parts.append('<div class="fresh-legend">当前显示 <b id="freshCount">%d</b> / %d 项｜点击上方卡片筛选</div>' % (summary['total'], summary['total']))
            parts.append('<div class="tablewrap"><table class="bal fresh-table"><tr><th>模块/指标</th><th>来源</th><th>频率</th><th>单位</th><th>最新标签/值</th><th>新鲜度</th><th>处理</th></tr>')
            for row in freshness['rows']:
                age = '，%d天' % row['age_days'] if row['age_days'] is not None else ''
                parts.append('<tr data-fresh="%s"><td><b>%s</b><br><span class="src-note">%s｜%s</span></td><td>%s</td><td>%s</td><td>%s</td><td>%s<br><span class="src-note">%s%s</span></td><td><span class="fresh-badge %s">%s</span></td><td>%s</td></tr>' % (
                    row['status'], shared.esc(row['name']), shared.esc(row['module']), shared.esc(row['key']), shared.esc(row['source']),
                    shared.esc(row['frequency']), shared.esc(row['unit']), shared.esc(row['latest']), shared.esc(row['latest_value']), shared.esc(age),
                    row['status'], shared.esc(row['status_label']), shared.esc(row['action'])))
            parts.append('</table></div>')
        else:
            parts.append('<div class="lead">%s · 最新读数逐项标注日期</div>' % shared.esc(title))
            parts.append(shared.chart_group(charts, keys))
        parts.append('</div></section>')

    payload = json.dumps(charts, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    doc = '''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow,noarchive"><link href="data:," rel="icon">
<title>研究周报 · 铝</title><style>__BASE_CSS__</style><link href="weekly_mobile.css" rel="stylesheet"><link href="weekly_polish.css" rel="stylesheet"><link href="research.css" rel="stylesheet"><link href="../weekly-light.css" rel="stylesheet"><style>.chart-box{height:240px;position:relative}.chart-box canvas{display:block;width:100%;height:100%}.tablewrap{overflow-x:auto}.page-body p{font-size:14.5px}.cbox .stats{line-height:1.55}.fresh-grid{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:9px;margin:12px 0}.fresh-kpi{border:1px solid var(--line);background:var(--cardbg);padding:10px;text-align:left;cursor:pointer}.fresh-kpi span,.fresh-kpi small{display:block;color:var(--faint);font-size:11px}.fresh-kpi strong{display:block;font-size:25px}.fresh-kpi.fresh strong{color:#188865}.fresh-kpi.watch strong{color:#a3760e}.fresh-kpi.stale strong,.fresh-kpi.missing strong{color:#c44b5d}.fresh-kpi.review strong{color:#7c4dff}.fresh-kpi.active{outline:2px solid var(--blue);outline-offset:1px}.fresh-legend{margin:12px 0;color:var(--grey);font-size:13px}.fresh-badge{display:inline-block;padding:1px 7px;border:1px solid var(--line);white-space:nowrap}.fresh-badge.fresh{color:#188865}.fresh-badge.watch{color:#a3760e}.fresh-badge.stale,.fresh-badge.missing{color:#c44b5d}.fresh-badge.review{color:#7c4dff}.src-note{color:var(--faint);font-size:11px}.fresh-table{min-width:1080px}.fresh-table td{vertical-align:top}@media(max-width:900px){.fresh-grid{grid-template-columns:repeat(3,1fr)}}</style></head>
<body><header class="research-topbar"><div class="research-brand"><a aria-label="返回整合研究站" class="back-home" href="../index.html">← 整合研究站</a><div class="brand-copy"><p class="research-eyebrow">COMMODITY WEEKLY RESEARCH · ALUMINUM</p><h1>研究周报 <span class="commodity-current">铝</span></h1><p>铝土矿 · 氧化铝 · 电解铝 · 全球库存 · 铝加工需求</p></div></div><div class="research-actions"><nav class="commodity-switch" aria-label="品种切换"><a href="../tin-weekly/#sumsec">锡</a><a href="../nickel-weekly/#sumsec">镍</a><a href="../zinc-weekly/#sumsec">锌</a><a class="active" href="#sumsec">铝</a><a href="../copper-weekly/#sumsec">铜</a><a href="../lead-weekly/#sumsec">铅</a></nav><a class="back-home" href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">铝企财报 ↗</a><span class="research-stamp">数据截至 <b>__ASOF__</b></span><button aria-pressed="false" class="view-toggle" id="viewToggle" type="button">切换手机版</button></div></header>
<div class="research-workspace"><aside class="research-sidebar top-navigation" id="researchSidebar"><button aria-controls="treeNav" aria-expanded="false" class="menu-toggle" id="menuToggle" type="button">研究目录 <span>＋</span></button><div class="sidebar-heading"><span>RESEARCH INDEX</span><b>研究目录</b></div><nav aria-label="铝研究分层目录" class="tree-nav top-nav-inner" id="treeNav"></nav></aside><div class="research-content"><div aria-live="polite" class="breadcrumb" id="breadcrumb"></div><div class="wrap" id="reportWrap">__SECTIONS__</div><section aria-live="polite" class="tool-page" hidden id="toolPage"><div class="tool-head"><div><p class="research-eyebrow">RESEARCH WORKBENCH</p><h2 id="toolTitle">研究工作台</h2></div><button id="closeTool" type="button">← 返回周报</button></div><iframe hidden id="toolFrame" title="铝研究工作台"></iframe></section></div></div>
<script src="lib/chart.min.js"></script><script>const charts=__PAYLOAD__;document.querySelectorAll('canvas[data-chart]').forEach(function(canvas){const c=charts[canvas.dataset.chart];if(!c)return;const ds=(c.datasets||[]).map(function(d){return Object.assign({borderWidth:1.5,pointRadius:0,spanGaps:true},d);});const cfg={type:'line',data:{labels:c.labels||[],datasets:ds},options:{responsive:true,maintainAspectRatio:false,animation:false,interaction:{mode:'index',intersect:false},plugins:{legend:{position:'bottom',labels:{boxWidth:12,font:{family:'KaiTi,"楷体",serif',size:10}}}},scales:{x:{ticks:{maxTicksLimit:8,font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}},y:{ticks:{font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}}}}};new Chart(canvas,cfg);});</script><script src="weekly_mobile.js"></script><script src="research.js"></script><script>document.querySelectorAll('.fresh-kpi').forEach(function(btn){btn.addEventListener('click',function(){const code=btn.dataset.filter;let shown=0;document.querySelectorAll('.fresh-kpi').forEach(function(x){x.classList.toggle('active',x===btn)});document.querySelectorAll('tr[data-fresh]').forEach(function(row){const visible=code==='all'||row.dataset.fresh===code;row.hidden=!visible;if(visible)shown++});document.getElementById('freshCount').textContent=shown})});</script></body></html>'''
    doc = doc.replace('__BASE_CSS__', base_css).replace('__ASOF__', shared.esc(asof)).replace('__SECTIONS__', ''.join(parts)).replace('__PAYLOAD__', payload)
    tmp = OUT / 'index.html.tmp'
    tmp.write_text(doc, encoding='utf-8')
    os.replace(tmp, OUT / 'index.html')
    summary = {'title': '研究周报 · 铝', 'asof': asof, 'headline': headline, 'metrics': metrics,
               'local_url': 'aluminum-weekly/#sumsec', 'source_url': 'https://amiya866.github.io/caibao-hub/',
               'charts': len(charts), 'freshness': freshness['summary'], 'sync_status': 'ok'}
    (BASE / 'data' / 'aluminum_weekly.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    update_daily(summary)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == '__main__':
    main()

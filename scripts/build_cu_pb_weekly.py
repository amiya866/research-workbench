# -*- coding: utf-8 -*-
"""用与锌/铝完全相同的结构生成铜、铅研究周报。"""
import csv
import datetime
import json
import os
import re
import shutil
import sqlite3
import subprocess
from pathlib import Path

from openpyxl import load_workbook

import build_nickel_weekly as shared

BASE = Path(__file__).resolve().parents[1]
CACHE = Path(r'D:/Kimi/金属总网/网站构建/db/数据缓存_v1.db')
SHFE_CSV = Path(r'D:/Kimi/有色截面配置_2026/_价格缓存.csv')
LME_CSV = Path(r'D:/Kimi/有色截面配置_2026/_LME价格缓存.csv')
IFIND_DB = Path(r'D:/Kimi/ifind_mirror/ifind_mirror.db')
BAICHUAN_DB = Path(r'D:/Kimi/baichuan_mirror/baichuan_mirror.db')
NW = getattr(subprocess, 'CREATE_NO_WINDOW', 0)
MONTHS = {'January': 1, 'February': 2, 'March': 3, 'April': 4, 'May': 5, 'June': 6,
          'July': 7, 'August': 8, 'September': 9, 'October': 10, 'November': 11, 'December': 12}

COMMON_SECTIONS = [
    ('sumsec', '本周总结', '总览', []),
    ('price', '价格与结构', '总览', []),
    ('supply_mine', '矿端与进口', '重点关注/供给', []),
    ('supply_smelt', '冶炼与精炼', '重点关注/供给', []),
    ('inventory', '显性库存', '重点关注/供给', []),
    ('demand', '需求与终端', '重点关注/需求', []),
    ('cost', '成本与利润', '平衡表', []),
    ('balance', '年度平衡与2027策略', '平衡表', []),
    ('framework', '产业框架', '总览', []),
    ('technical', '技术分析', '宏观与策略', ['overview_shfe']),
    ('sources', '数据缺口', '总览', []),
]

CONFIG = {
    'CU': {
        'cn': '铜', 'en': 'COPPER', 'out': 'copper-weekly', 'prefix': 'copper',
        'subtitle': '铜精矿TC · 冶炼利润 · 全球库存 · 电线电缆 · 成本平衡',
        'lme_field': 'LME_Cu_cash',
        'domestic': ('cache', 'ID01677008', 'SMM电解铜升贴水', '元/吨'),
        'financial': Path(r'D:/Kimi/金属总网/网站构建/财报汇总/铜/全球铜企季度产量梳理.xlsx'),
        'sheets': [('铜矿·季度产量', 12), ('国内铜企·季度产量', 10)],
        'price': ['overview_shfe', 'overview_lme', 'overview_lme03', 'overview_domestic_premium', 's20015675', 'ID01030232', 'FU00017551'],
        'supply_mine': ['ID01154994', 'a10015145', 'a10015126', 'a10015120', 'a10015111', 'a10015108', 'ID01002000'],
        'supply_smelt': ['j02870870', 'a12855621', 'ID01137328', 'ID01137352'],
        'inventory': ['overview_global_stock', 'FU00015680', 'a10134435', 'a10083949', 'ID00188319', 'FU00015497', 'FU00015498', 'ID00188286'],
        'demand': ['ID01067717', 'a10015025', 'a10015276', 'a10014968', 'CM0000753806', 'a10015356'],
        'cost': ['ID01241097', 'ID01723352'],
        'kpi': [('shfe', '沪铜主连', '元/吨'), ('lme', 'LME铜', '美元/吨'), ('global_stock', 'LME+SHFE库存', '吨'),
                ('social_stock', '电解铜社库', '万吨'), ('tc', '铜精矿TC', '美元/吨'), ('import_pnl', '电解铜进口盈亏', '元/吨'),
                ('lme03', 'LME 0-3', '美元/吨'), ('domestic', '国内现货升贴水', '元/吨')],
        'latest': {'social_stock': 'a10083949', 'tc': 'ID01154994', 'import_pnl': 'ID01030232',
                   'lme03': 'overview_lme03', 'domestic': 'overview_domestic_premium'},
        'demand_key': 'ID01067717', 'output_key': 'j02870870',
        'js_supply': ['铜矿与进口', '冶炼与精炼', '显性库存', '成本与平衡'],
        'js_demand': '铜杆/铜管/电线电缆',
    },
    'PB': {
        'cn': '铅', 'en': 'LEAD', 'out': 'lead-weekly', 'prefix': 'lead',
        'subtitle': '铅精矿TC · 原生/再生铅 · 全球库存 · 铅酸电池 · 成本平衡',
        'lme_field': 'LME_Pb_cash',
        'domestic': ('ifind', 'S005068211', '1#铅锭升贴水（华东）', '元/吨'),
        'financial': Path(r'D:/Kimi/金属总网/网站构建/财报汇总/铅/全球铅企季度产量梳理.xlsx'),
        'sheets': [('铅精矿·季度产量', 12), ('精炼铅·季度产量', 10)],
        'price': ['overview_shfe', 'overview_lme', 'overview_lme03', 'overview_domestic_premium', 'overview_domestic_premium_baichuan'],
        'supply_mine': ['a10127385', 'a10017055', 'a10021529', 'a10021532', 'a10021534', 'a10021525'],
        'supply_smelt': ['ID01001562', 'a10016988', 'a10017062', 'a10017000', 'a10017037', 'a10017091'],
        'inventory': ['overview_global_stock', 'FU00015325', 'a10134441', 'ID00188315'],
        'demand': ['ID01030005'],
        'cost': ['a10127388', 'ID01245255'],
        'kpi': [('shfe', '沪铅主连', '元/吨'), ('lme', 'LME铅', '美元/吨'), ('global_stock', 'LME+SHFE库存', '吨'),
                ('social_stock', '铅锭社库', '万吨'), ('tc', '铅精矿TC', '元/金属吨'), ('output', '原生铅产量', '万吨'),
                ('lme03', 'LME 0-3', '美元/吨'), ('domestic', '国内升贴水（百川）', '元/吨')],
        'latest': {'social_stock': 'ID00188315', 'tc': 'a10127385', 'output': 'ID01001562',
                   'lme03': 'overview_lme03', 'domestic': 'overview_domestic_premium_baichuan'},
        'demand_key': 'ID01030005', 'output_key': 'ID01001562',
        'js_supply': ['铅矿与进口', '原生/再生铅', '显性库存', '成本与平衡'],
        'js_demand': '铅酸蓄电池',
    },
}

BALANCE_CONFIGS = {
    'CU': {
        'title': '年度平衡与2027策略：精炼过剩扩大，但矿端长期紧张仍在',
        'rows': [
            ('2026E', '精炼产量+0.4%', '精炼消费+1.5%', '+10.4', '精炼小幅过剩，矿端仍紧'),
            ('2027E', '精炼产量+3.4%', '精炼消费+2.5%', '+36.8', '过剩扩大，远月压力增加'),
        ],
        'regional': '2027精炼恢复并不等于矿山长期宽松。Grasberg、Cobre Panamá、OT与冶炼/再生恢复使精炼增速超过需求；ICSG同时预计2027矿山增速3.1%，但2028以后仍需新增矿投资。',
        'source': '来源：ICSG 2026-10-14；2026/2027矿山产量增速分别1.3%/3.1%。平衡为表观精炼口径，未计中国非报告库存变化。',
        'strategy': [
            '方向：不照搬铝的长期空配。2027精炼过剩限制远月高度，但矿端扰动和长期需求仍支持回调找买点。',
            '入场：仅在LME/国内库存同步上升、Back与升水走弱、矿山和冶炼恢复共振时做阶段空配；不把低TC直接当成做空铜锭的信号。',
            '风控：若Grasberg/Cobre Panamá恢复不及预期、TC继续下行或全球库存重新下降，停止空配；长期核心仍是矿端约束。',
            '止盈/失效：阶段回落4%—8%兑现；若2027精炼过剩被证伪，恢复回调偏多而非逆势扛单。',
        ],
        'drivers': [
            '恢复：Grasberg爬产、Cobre Panamá三条产线恢复、Oyu Tolgoi与Batu Hijau，叠加印尼/印度/欧盟冶炼爬产和再生铜。',
            '约束：2026精炼仅+0.4%，说明矿端紧张不会立即变成铜锭过剩；2027才更明显。',
            '需求：电网、能源转型和AI基础设施仍增长，但2026消费仅+1.5%，需要用加工开工和库存验证，而非只看长期叙事。',
        ],
    },
    'PB': {
        'title': '年度平衡与2027策略：再生供应恢复，过剩扩大，中期空配较明确',
        'rows': [
            ('2026E·矿端', '矿山444万吨（-2.8%）', '—', '矿端收缩', '原生矿提供阶段支撑'),
            ('2026E·精炼', '1,388万吨（+1.0%）', '1,379万吨（+0.9%）', '+9.0', '精炼已过剩'),
            ('2027E·精炼', '1,408万吨（+1.5%）', '1,392万吨（+1.0%）', '+16.3', '过剩扩大，空配优先级提高'),
        ],
        'regional': '铅的核心不是单看矿山，而是再生铅与新增二次产能。2027矿山恢复至458万吨（+2.9%），精炼增速继续高于需求；中国需求仅小幅增长且2027预期下降。',
        'source': '来源：ILZSG 2026-10-19。铅平衡为精炼口径；再生铅新增产能和恢复是主要供给增量，预测不是实际值。',
        'strategy': [
            '方向：2027远月反弹空比锌更明确，但近端去库时不追空；优先等待再生铅利润、开工与库存共同转弱。',
            '入场：远月反弹且国内贴水扩大、再生铅开工恢复、社库止降时建立空配；铅酸电池旺季临近时降低仓位。',
            '风控：若原生+再生同时扰动、废电瓶价格令再生持续减产或库存连续两月下降，停止空配。',
            '止盈/失效：回落5%—8%分批兑现；若2027过剩收窄至平衡附近，退出长期空配。每月复核再生链。',
        ],
        'drivers': [
            '再生：印度、韩国、土耳其、巴西新增二次产能，2027中国精炼恢复，是过剩扩大的主因。',
            '原生：Garpenberg地震扰动、Antamina和其他矿山下降构成近端支撑；2027澳洲新产能带来恢复。',
            '需求：替换电池稳定，但铅酸出口和汽车产量限制中国需求；不能用铅价反推真实订单。',
        ],
    },
}


def series(conn, ind_id, name, unit):
    rows = conn.execute('select date,value from series_points where ind_id=? order by date', (ind_id,)).fetchall()
    return {'labels': [r[0] for r in rows], 'datasets': [{'label': name, 'data': [r[1] for r in rows],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}, unit


def csv_chart(path, column, name):
    labels, values = [], []
    with open(path, encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            value = row.get(column)
            if value not in (None, ''):
                labels.append(row[''])
                values.append(float(value))
    return {'labels': labels, 'datasets': [{'label': name, 'data': values, 'borderColor': '#1F4E79',
            'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def lme03_chart(field, name):
    labels, values = [], []
    today = datetime.date.today()
    for year in (today.year - 1, today.year):
        url = 'https://www.westmetall.com/en/markdaten.php?action=table&field=%s&year=%d' % (field, year)
        try:
            text = subprocess.run(['curl', '-sL', '--max-time', '30', '-H', 'User-Agent: Mozilla/5.0', url],
                                  capture_output=True, timeout=45, creationflags=NW).stdout.decode('utf-8', 'replace')
        except Exception:
            continue
        for row in re.findall(r'<tr[^>]*>(.*?)</tr>', text, re.S):
            cells = [re.sub(r'<[^>]+>|\s+', ' ', c).strip() for c in re.findall(r'<td[^>]*>(.*?)</td>', row, re.S)]
            if len(cells) < 3:
                continue
            match = re.match(r'(\d{2})\. (\w+) (\d{4})', cells[0])
            if not match or match.group(2) not in MONTHS:
                continue
            try:
                basis = float(cells[1].replace(',', '')) - float(cells[2].replace(',', ''))
            except ValueError:
                continue
            day = '%s-%02d-%02d' % (match.group(3), MONTHS[match.group(2)], int(match.group(1)))
            if day not in dict(zip(labels, values)):
                labels.append(day); values.append(basis)
    data = sorted(zip(labels, values))
    return {'labels': [x[0] for x in data], 'datasets': [{'label': name, 'data': [x[1] for x in data],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def ifind_series(code, name, unit):
    if not IFIND_DB.exists():
        return None
    conn = sqlite3.connect(f'file:{IFIND_DB.as_posix()}?mode=ro', uri=True)
    rows = conn.execute('select date,value from series_points where code=? order by date', (code,)).fetchall()
    conn.close()
    points = []
    for date, value in rows:
        try:
            points.append((date, float(value)))
        except (TypeError, ValueError):
            continue
    if not points:
        return None
    return {'labels': [r[0] for r in points], 'datasets': [{'label': name, 'data': [r[1] for r in points],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}, unit


def baichuan_series(key, name, unit):
    if not BAICHUAN_DB.exists():
        return None
    conn = sqlite3.connect(f'file:{BAICHUAN_DB.as_posix()}?mode=ro', uri=True)
    rows = conn.execute('select date,value_num from series_points where indicator_key=? order by date', (key,)).fetchall()
    conn.close()
    points = [(date, float(value)) for date, value in rows if value is not None]
    if not points:
        return None
    return {'labels': [row[0] for row in points], 'datasets': [{'label': name, 'data': [row[1] for row in points],
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}, unit


def last(chart):
    point = shared.last_point(chart)
    return [point[0], point[1]] if point else [None, None]


def global_stock(lme, shfe):
    dates = sorted(set(lme['labels']) | set(shfe['labels']))
    lm = dict(zip(lme['labels'], lme['datasets'][0]['data']))
    sm = dict(zip(shfe['labels'], shfe['datasets'][0]['data']))
    values, last_l, last_s = [], None, None
    for date in dates:
        last_l = lm.get(date, last_l)
        last_s = sm.get(date, last_s)
        values.append(last_l + last_s if last_l is not None and last_s is not None else None)
    return {'labels': dates, 'datasets': [{'label': 'LME+SHFE库存', 'data': values,
            'borderColor': '#1F4E79', 'backgroundColor': '#1F4E79', 'borderWidth': 1.6}], '_type': 'line'}


def rsi(values, period=14):
    if len(values) <= period:
        return None
    changes = [b - a for a, b in zip(values[-period - 1:-1], values[-period:])]
    gains = sum(max(x, 0) for x in changes) / period
    losses = sum(max(-x, 0) for x in changes) / period
    return 100 if not losses else 100 - 100 / (1 + gains / losses)


def ema(values, period):
    alpha = 2 / (period + 1)
    out = []
    for i, value in enumerate(values):
        out.append(value if i == 0 else alpha * value + (1 - alpha) * out[-1])
    return out


def update_daily(cfg, summary):
    path = BASE / 'reports' / 'daily_latest.html'
    if not path.exists():
        return
    page = path.read_text(encoding='utf-8')
    begin = '<!-- %s_WEEKLY_CARD_BEGIN -->' % cfg['en']
    end = '<!-- %s_WEEKLY_CARD_END -->' % cfg['en']
    metrics = ''.join('<li><b>%s</b>　%s</li>' % (shared.esc(m['name']), shared.esc(m['value'])) for m in summary['metrics'][:4])
    card = '''%s
<div class="card" id="%sWeeklyCard" style="border-left:3px solid var(--gold);margin:14px 0;">
  <div style="display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap;align-items:center;">
    <h3 style="margin:0;color:var(--gold);">研究周报 · %s</h3>
    <a href="../%s/#sumsec" style="font-weight:700;">打开完整%s周报 ↗</a>
  </div>
  <p style="margin-top:8px;">%s</p>
  <ul style="margin:8px 0 0 18px;font-size:12px;color:var(--dim);">%s</ul>
  <p style="font-size:11px;color:var(--dim);margin-top:8px;">数据截至 %s · 本地生成，随研究站管线自动更新 · <a href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">%s企财报</a></p>
</div>
%s''' % (begin, cfg['prefix'], cfg['cn'], cfg['out'], cfg['cn'], shared.esc(summary['headline']), metrics,
           shared.esc(summary['asof']), cfg['cn'], end)
    pattern = re.escape(begin) + r'.*?' + re.escape(end)
    if re.search(pattern, page, re.S):
        page = re.sub(pattern, lambda _m: card, page, count=1, flags=re.S)
    else:
        anchor = '<!-- NICKEL_WEEKLY_CARD_END -->'
        if anchor not in page:
            anchor = '<section class="dp on" id="dp1">'
        page = page.replace(anchor, anchor + '\n' + card, 1)
    path.write_text(page, encoding='utf-8')


def make_outlook(metal, q3, q4, latest, demand, output):
    if metal == 'CU':
        return '''<details class="fold" id="q4outlook"><summary>季度展望：Q3复盘与Q4展望（点击展开/折叠）</summary><div class="foldbody">
<div class="lead">一、三季度复盘与四季度展望</div><ul class="sq">
<li>Q3沪铜价格变动 %s，矿端TC低位、宏观利率与库存共同定价，铜价对供应扰动和美元变化保持较高敏感度。</li>
<li>Q4迄今价格变动 %s。当前LME与SHFE合计库存 %s 吨、国内社库 %s 万吨，库存仍是判断近月结构的第一约束。</li>
<li>结构定位：LME 0-3 %s 美元/吨（%s），国内现货升贴水 %s 元/吨（%s）；期现结构必须与库存方向一致才有效。</li>
<li>基准判断是“矿端偏紧提供底部支撑，宏观与需求决定上方弹性”。不能把低TC直接等同于精炼铜即时短缺，需等待产量、库存和现货共同验证。</li></ul>
<div class="lead">二、供应：TC与冶炼利润共同判断</div><ul class="sq">
<li>铜精矿TC最新 %s 美元/吨（%s），低TC反映矿端议价偏紧，但冶炼厂还受硫酸收益、检修和原料库存影响。</li>
<li>冶炼现货利润最新 %s 元/吨（%s），Q4重点看低TC能否传导为实际减产，而不是只停留在报价层面。</li>
<li>海外矿山Q3财报与爬产节奏提供边际增量，仍需核查运输、冶炼和仓单交付。</li></ul>
<div class="lead">三、需求与库存：电网托底，加工开工验证</div><ul class="sq">
<li>精铜制杆产能利用率最新 %s%%（%s），电线电缆、铜管、铜板带和电网投资共同决定真实消费。</li>
<li>若社库下降且升贴水、制杆开工同步改善，近月结构更有支撑；若高价格抑制补库，库存拐点会先削弱远月估值。</li>
<li>需求端需区分电网投资、家电和新能源链条，不能只用铜价反推景气。</li></ul>
<div class="lead">四、策略与验证点</div><ul class="sq">
<li>近月看库存、升贴水与进口盈亏，远月看矿山爬产和宏观利率，避免在低库存环境下追逐极端行情。</li>
<li>上修条件：TC继续低位＋持续去库＋加工开工改善；下修条件：冶炼恢复＋仓单回升＋终端补库走弱。</li>
<li>本段为数据驱动的研究框架，不构成投资建议。</li></ul></div></details>''' % (
            shared.pct(q3), shared.pct(q4), shared.fmt(latest['global_stock'][1]), shared.fmt(latest['social_stock'][1]),
            shared.fmt(latest['lme03'][1]), latest['lme03'][0], shared.fmt(latest['domestic'][1]), latest['domestic'][0],
            shared.fmt(latest['tc'][1]), latest['tc'][0], shared.fmt(output[1]) if output else '—', output[0] if output else '—',
            shared.fmt(demand[1]) if demand else '—', demand[0] if demand else '—')
    return '''<details class="fold" id="q4outlook"><summary>季度展望：Q3复盘与Q4展望（点击展开/折叠）</summary><div class="foldbody">
<div class="lead">一、三季度复盘与四季度展望</div><ul class="sq">
<li>Q3沪铅价格变动 %s，原生铅、再生铅和铅精矿TC共同塑造供应弹性，铅酸电池需求决定价格上方空间。</li>
<li>Q4迄今价格变动 %s。当前LME与SHFE合计库存 %s 吨、国内社库 %s 万吨，库存与再生铅利润是判断近月压力的重要证据。</li>
<li>结构定位：LME 0-3 %s 美元/吨（%s），国内现货升贴水 %s 元/吨（%s）；期现结构必须与再生供应和库存方向一致才有效。</li>
<li>基准判断是“矿端与再生成本提供支撑，供应恢复约束趋势高度”。铅不能只看原生矿TC，还需同时跟踪废电瓶和再生产量。</li></ul>
<div class="lead">二、供应：原生与再生两条线</div><ul class="sq">
<li>国产铅精矿TC最新 %s 元/金属吨（%s），需结合进口矿、原生铅开工和精炼产量判断矿端松紧。</li>
<li>原生铅最新产量 %s 万吨（%s），再生粗铅与废电瓶价格决定另一条供应曲线，不能相互替代。</li>
<li>若再生利润恢复带动开工上行，铅锭供应可能比原生矿端更快改善。</li></ul>
<div class="lead">三、需求与库存：铅酸电池是核心验证</div><ul class="sq">
<li>铅酸蓄电池开工率最新 %s%%（%s），汽车启动、储能和替换电池共同决定真实订单。</li>
<li>若电池开工稳定且社库下降，近月结构仍有支撑；若再生产量恢复并带来累库，远月压力会先体现。</li>
<li>需求端不能用铅价反推，必须观察电池开工、出口和终端替换需求。</li></ul>
<div class="lead">四、策略与验证点</div><ul class="sq">
<li>近月看库存与现货结构，远月看再生供应恢复，避免仅凭TC或废电瓶价格单边外推。</li>
<li>上修条件：持续去库＋电池开工稳定＋原生/再生供应扰动；下修条件：再生恢复超预期＋仓单回升＋电池开工走弱。</li>
<li>本段为数据驱动的研究框架，不构成投资建议。</li></ul></div></details>''' % (
        shared.pct(q3), shared.pct(q4), shared.fmt(latest['global_stock'][1]), shared.fmt(latest['social_stock'][1]),
        shared.fmt(latest['lme03'][1]), latest['lme03'][0], shared.fmt(latest['domestic'][1]), latest['domestic'][0],
        shared.fmt(latest['tc'][1]), latest['tc'][0], shared.fmt(output[1]) if output else '—', output[0] if output else '—',
        shared.fmt(demand[1]) if demand else '—', demand[0] if demand else '—')


def sign(value):
    return 1 if (value or 0) > 0 else (-1 if (value or 0) < 0 else 0)


def marginal_items_for(metal, charts, latest, chg5, structure_desc, demand, output):
    price_score = sign(chg5)
    price_judge = '价格上行，偏多证据' if price_score > 0 else ('价格下行，偏空证据' if price_score < 0 else '价格震荡，等待方向')
    lme_score = sign(latest['lme03'][1])
    domestic_score = sign(latest['domestic'][1])
    if lme_score and domestic_score and lme_score == domestic_score:
        structure_score = lme_score * 2
        structure_judge = '内外结构同向支撑（方向较明确）' if structure_score > 0 else '内外结构同向偏弱（方向较明确）'
    elif lme_score or domestic_score:
        structure_score = lme_score or domestic_score
        structure_judge = '单边结构改善，仍需另一边确认' if structure_score > 0 else '单边结构转弱，仍需另一边确认'
    else:
        structure_score = 0
        structure_judge = '结构数据不足，跟踪'
    global_chg = shared.change(charts['overview_global_stock'], 20)
    social_key = 'a10083949' if metal == 'CU' else 'ID00188315'
    social_chg = shared.change(charts[social_key], 20)
    if global_chg is not None and social_chg is not None and sign(global_chg) == sign(social_chg):
        inventory_score = -sign(global_chg)
        inventory_judge = '内外库存同向去化，近月有支撑' if inventory_score > 0 else '内外库存同向累积，近月承压'
    elif global_chg is not None or social_chg is not None:
        inventory_score = -sign(global_chg if global_chg is not None else social_chg)
        inventory_judge = '单边库存方向成立，仍需另一边确认' if inventory_score > 0 else '单边库存方向偏弱，仍需另一边确认'
    else:
        inventory_score = 0
        inventory_judge = '库存变化不可得，跟踪'
    if metal == 'CU':
        tc_chg = shared.change(charts['ID01154994'], 5)
        profit_chg = shared.change(charts['j02870870'], 20)
        supply_evidence = '铜精矿TC %s 美元/吨（近5日 %s），冶炼现货利润近20日 %s' % (
            shared.fmt(latest['tc'][1]), shared.pct(tc_chg), shared.pct(profit_chg))
        if latest['tc'][1] < -50 and (tc_chg or 0) < 0:
            supply_score, supply_judge = 1, 'TC低位且继续恶化，矿端偏紧（方向明确）'
        elif (tc_chg or 0) > 0.1:
            supply_score, supply_judge = -1, 'TC低位回升，矿端边际转松'
        else:
            supply_score, supply_judge = 0, '矿端仍紧但边际未继续恶化，跟踪'
    else:
        output_chg = shared.change(charts['ID01001562'], 20)
        supply_evidence = '原生铅产量 %s 万吨（%s，近月变动 %s）' % (
            shared.fmt(output[1]) if output else '—', output[0] if output else '—', shared.pct(output_chg))
        if (output_chg or 0) < -0.02:
            supply_score, supply_judge = 1, '原生供应收缩，供给端有支撑'
        elif (output_chg or 0) > 0.02:
            supply_score, supply_judge = -1, '原生供应恢复，供给端施压'
        else:
            supply_score, supply_judge = 0, '原生供应变化不显著，跟踪再生线'
    demand_chg = shared.change(charts['ID01067717' if metal == 'CU' else 'ID01030005'], 20)
    demand_evidence = '%s %s（20日变动 %s）' % (
        '精铜制杆产能利用率' if metal == 'CU' else '铅酸蓄电池开工率',
        shared.fmt(demand[1]) if demand else '—', shared.pct(demand_chg))
    if (demand_chg or 0) > 0.02:
        demand_score, demand_judge = 1, '需求读数改善，偏多证据'
    elif (demand_chg or 0) < -0.02:
        demand_score, demand_judge = -1, '需求读数走弱，偏空证据'
    else:
        demand_score, demand_judge = 0, '需求变化不足，跟踪'
    return [
        ('价格', '近5个交易日 %s' % shared.pct(chg5), price_judge, price_score),
        ('结构', structure_desc, structure_judge, structure_score),
        ('库存', '全球显性库存20日 %s，国内社库20日 %s' % (shared.pct(global_chg), shared.pct(social_chg)), inventory_judge, inventory_score),
        ('供应', supply_evidence, supply_judge, supply_score),
        ('需求', demand_evidence, demand_judge, demand_score),
    ]


def main(metal):
    cfg = CONFIG[metal]
    out = BASE / cfg['out']
    sections = []
    for sid, title, trail, keys in COMMON_SECTIONS:
        sections.append((sid, title, trail, list(keys) if keys else list(cfg.get(sid, []))))
    conn = sqlite3.connect(f'file:{CACHE.as_posix()}?mode=ro', uri=True)
    charts, source_meta, names, units = {}, {}, {}, {'overview_shfe': '元/吨', 'overview_lme': '美元/吨', 'overview_global_stock': '吨'}
    for ind_id, name, unit in conn.execute('select ind_id,name,unit from series_meta where comm=?', (metal,)):
        charts[ind_id], units[ind_id] = series(conn, ind_id, name, unit)
        source_meta[ind_id] = {'type': 'zhiji', 'indicator': ind_id, 'unit': unit}
        names[ind_id] = name
    conn.close()
    cn = cfg['cn']
    charts['overview_shfe'] = csv_chart(SHFE_CSV, metal, '沪%s主连' % cn)
    charts['overview_lme'] = csv_chart(LME_CSV, metal, 'LME%s' % cn)
    if cfg.get('lme_field'):
        charts['overview_lme03'] = lme03_chart(cfg['lme_field'], 'LME %s 0-3 升贴水' % cn)
        source_meta['overview_lme03'] = {'type': 'westmetall', 'formula': 'cash-3M'}
        names['overview_lme03'] = 'LME %s 0-3 升贴水' % cn
        units['overview_lme03'] = '美元/吨'
    domestic_kind, domestic_id, domestic_name, domestic_unit = cfg['domestic']
    if domestic_kind == 'ifind':
        domestic_result = ifind_series(domestic_id, domestic_name, domestic_unit)
        if domestic_result:
            charts['overview_domestic_premium'], units['overview_domestic_premium'] = domestic_result
            source_meta['overview_domestic_premium'] = {'type': 'iFinD EDB', 'indicator': domestic_id}
        if metal == 'PB':
            baichuan_name = '华东1#铅锭现货升贴水（百川口径）'
            baichuan_result = baichuan_series('1588348470396512053', baichuan_name, domestic_unit)
            if baichuan_result:
                charts['overview_domestic_premium_baichuan'], units['overview_domestic_premium_baichuan'] = baichuan_result
                names['overview_domestic_premium_baichuan'] = baichuan_name
                source_meta['overview_domestic_premium_baichuan'] = {
                    'type': '百川盈孚', 'indicator': 'price:1588348470396512053',
                    'note': '独立并列序列；与iFinD 99.994%华东均价口径不同，不拼接'}
    else:
        charts['overview_domestic_premium'] = json.loads(json.dumps(charts[domestic_id]))
        units['overview_domestic_premium'] = domestic_unit
        source_meta['overview_domestic_premium'] = {'type': 'zhiji', 'indicator': domestic_id}
    names['overview_domestic_premium'] = domestic_name
    lme_key = 'a10134435' if metal == 'CU' else 'a10134441'
    shfe_key = 'FU00015680' if metal == 'CU' else 'FU00015325'
    charts['overview_global_stock'] = global_stock(charts[lme_key], charts[shfe_key])
    for key, label in [('overview_shfe', 'guan+本地日K'), ('overview_lme', '新浪外盘+本地缓存'), ('overview_global_stock', 'LME+SHFE合计')]:
        source_meta[key] = {'type': 'derived', 'formula': label}
    names.update({'overview_shfe': '沪%s主连' % cn, 'overview_lme': 'LME%s' % cn, 'overview_global_stock': '全球显性库存'})
    shared.NAMES = names
    shared.SECTIONS = sections
    built = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
    data = {'charts': charts, 'sourceMeta': {'charts': source_meta}, 'meta': {'builtAt': built, 'warnings': [], 'emptyCharts': []}}
    freshness = shared.build_freshness(data)
    (BASE / 'data' / ('%s_weekly_freshness.json' % cfg['prefix'])).write_text(json.dumps(freshness, ensure_ascii=False, indent=2), encoding='utf-8')
    latest = {'shfe': last(charts['overview_shfe']), 'lme': last(charts['overview_lme']),
              'global_stock': last(charts['overview_global_stock'])}
    for key, ind_id in cfg['latest'].items():
        latest[key] = last(charts[ind_id])
    asof = max(value[0] for value in latest.values() if value[0])
    shfe_values = [v for v in charts['overview_shfe']['datasets'][0]['data'] if v is not None]
    chg5 = shfe_values[-1] / shfe_values[-6] - 1 if len(shfe_values) > 6 else None
    stock_values = [v for v in charts[lme_key]['datasets'][0]['data'] if v is not None]
    stock_chg = stock_values[-1] / stock_values[-6] - 1 if len(stock_values) > 6 else None
    ma5 = sum(shfe_values[-5:]) / 5
    ma20 = sum(shfe_values[-20:]) / 20
    current_rsi = rsi(shfe_values)
    macd = [a - b for a, b in zip(ema(shfe_values, 12), ema(shfe_values, 26))]
    signal = ema(macd, 9)
    price_word = '偏强' if (chg5 or 0) > 0.01 else ('偏弱' if (chg5 or 0) < -0.01 else '震荡')
    stock_word = '累库' if (stock_chg or 0) > 0 else '去库'
    lme_basis = latest['lme03'][1] or 0
    domestic_premium = latest['domestic'][1] or 0
    lme_structure = 'Back' if lme_basis > 0 else ('Contango' if lme_basis < 0 else '持平')
    domestic_structure = '升水' if domestic_premium > 0 else ('贴水' if domestic_premium < 0 else '持平')
    structure_word = '结构偏Back' if lme_basis > 0 else ('结构偏Contango' if lme_basis < 0 else '结构中性')
    structure_desc = 'LME 0-3 %s 美元/吨（%s），国内现货%s %s 元/吨（%s）' % (
        shared.fmt(latest['lme03'][1]), latest['lme03'][0], domestic_structure,
        shared.fmt(latest['domestic'][1]), latest['domestic'][0])
    if metal == 'CU':
        headline = ('本周定调：沪铜%s与LME%s并存，%s —— 沪铜近5个交易日 %s，LME库存近5个交易日 %s；'
                    '电解铜社库 %s 万吨（%s），铜精矿TC %s 美元/吨（%s）；结构：%s。'
                    '近端重点跟踪价格、结构、库存和供应能否同向验证；⚙️规则摘要，不构成投资建议。') % (
            price_word, stock_word, structure_word, shared.pct(chg5), shared.pct(stock_chg), shared.fmt(latest['social_stock'][1]), latest['social_stock'][0],
            shared.fmt(latest['tc'][1]), latest['tc'][0], structure_desc)
    else:
        headline = ('本周定调：沪铅%s与LME%s并存，%s —— 沪铅近5个交易日 %s，LME库存近5个交易日 %s；'
                    '铅锭社库 %s 万吨（%s），国产铅精矿TC %s 元/金属吨（%s）；结构：%s。'
                    '近端重点跟踪价格、结构、库存和供应能否同向验证；⚙️规则摘要，不构成投资建议。') % (
            price_word, stock_word, structure_word, shared.pct(chg5), shared.pct(stock_chg), shared.fmt(latest['social_stock'][1]), latest['social_stock'][0],
            shared.fmt(latest['tc'][1]), latest['tc'][0], structure_desc)
    metrics = [{'name': name, 'value': '%s %s（%s）' % (shared.fmt(latest[key][1]), unit, latest[key][0])}
               for key, name, unit in cfg['kpi']]
    q3 = shared.period_return(charts['overview_shfe'], '2026-07-01', '2026-09-30')
    q4 = shared.period_return(charts['overview_shfe'], '2026-10-01', asof)
    demand = shared.last_point(charts[cfg['demand_key']])
    output = shared.last_point(charts[cfg['output_key']])
    q4_html = shared.quarterly_html(out / 'index.html', make_outlook(metal, q3, q4, latest, demand, output))
    marginal_items = marginal_items_for(metal, charts, latest, chg5, structure_desc, demand, output)
    marginal_html = shared.marginal_block(marginal_items)
    social_trend = shared.change(charts[cfg['latest']['social_stock']], 20)
    if metal == 'CU':
        weekly_html = shared.weekly_outlook_html(
            '沪铜近5个交易日%s，%s；LME库存近5日%s、国内社库20日%s。近端Back与国内升水有支撑，但海外累库要求保留验证。' % (shared.pct(chg5), structure_desc, shared.pct(stock_chg), shared.pct(social_trend)),
            '中期至2027年，矿端偏紧与低TC仍是底部支撑，电网、能源转型和AI电力投资提供需求韧性；但矿山爬产、再生铜和冶炼产量可缓冲精炼端短缺，不能把低TC直接等同为铜锭立即短缺。',
            '中期保持回调偏多研究思路，近端不追高升水与Back；优先等海外库存止升、国内升水稳定和加工开工改善共同确认。若库存持续累积且结构转弱，先降低单边敞口。',
            '上修：矿山扰动/低TC持续＋全球库存下降＋加工开工改善。下修：矿山爬产和冶炼恢复＋LME与国内同步累库＋现货升水回落＋终端订单走弱。')
    else:
        weekly_html = shared.weekly_outlook_html(
            '沪铅近5个交易日%s，%s；LME库存近5日%s、国内社库20日%s。近端去库有支撑，但国内外现货结构仍偏弱。' % (shared.pct(chg5), structure_desc, shared.pct(stock_chg), shared.pct(social_trend)),
            '中期至2027年，铅酸电池需求成熟，再生铅供应对利润和废电瓶价格反应较快，决定铅的供应弹性通常高于铜锡；原生矿TC只能解释一部分供给，趋势性短缺需要原生与再生同时受限。',
            '以区间和结构策略优先，近端去库时不追空，反弹后重点观察再生利润、开工和电池排产；若再生恢复并重新累库，逐步提高偏空或卖出套保权重。',
            '上修：原生与再生同时受限＋电池开工稳定＋持续去库。下修：再生利润和产量恢复＋社库/仓单回升＋电池开工及替换需求走弱。')
    marginal_scores = {item[0]: item[3] for item in marginal_items}
    lme03_trend = shared.change(charts['overview_lme03'], 20)
    domestic_trend = shared.change(charts['overview_domestic_premium'], 20)
    social_trend = shared.change(charts[cfg['latest']['social_stock']], 20)
    charts = shared.seasonalize_charts(charts)
    out.mkdir(parents=True, exist_ok=True)
    (out / 'lib').mkdir(exist_ok=True)
    for name in ('research.css', 'weekly_mobile.css', 'weekly_polish.css', 'weekly_mobile.js'):
        shutil.copyfile(BASE / 'nickel-weekly' / name, out / name)
    shutil.copyfile(BASE / 'nickel-weekly' / 'lib' / 'chart.min.js', out / 'lib' / 'chart.min.js')
    shutil.copyfile(BASE / 'scripts' / 'cu_pb_weekly_research.js', out / 'research.js')
    base_css = re.search(r'<style>(.*?)</style>', (BASE / 'nickel-weekly' / 'index.html').read_text(encoding='utf-8'), re.S).group(1)

    parts = []
    for sid, title, trail, keys in sections:
        parts.append('<section class="report-page" data-page="%s" data-trail="%s"><div class="page-heading"><h2>%s</h2><span>%s</span></div><div class="page-body">' % (
            sid, shared.esc(trail), shared.esc(title), shared.esc(trail.replace('/', ' / ').upper())))
        if sid in ('sumsec', 'balance'):
            parts.append(shared.forecast_balance_html(BALANCE_CONFIGS[metal], shfe_values, full=(sid == 'balance')))
        if sid == 'sumsec':
            parts.append('<div class="tone"><b>%s</b></div><div class="roll">策略每日滚动更新：周报为周五存档，最新决策见日报页。使用左侧目录切换研究内容。</div>' % shared.esc(headline))
            parts.append(weekly_html)
            parts.append(q4_html)
            parts.append('<div id="sumsec"><h2>总结<span class="en">SUMMARY</span></h2><table class="sum"><tr><th>环节</th><th>相关数据</th><th>周度结论</th><th>跟踪</th></tr>')
            def _cls(score):
                return 'up' if (score or 0) > 0 else ('down' if (score or 0) < 0 else 'nt')
            lme_conclusion = lme_structure
            lme_status = 'Back偏多' if lme_basis > 0 else ('Contango偏空' if lme_basis < 0 else '中性')
            lme_cls = 'up' if lme_basis > 0 else ('down' if lme_basis < 0 else 'nt')
            if lme_basis > 0 and lme03_trend is not None and lme03_trend < 0:
                lme_conclusion = 'Back仍在但20日走弱，继续跟踪是否转Contango'; lme_status = '关注'; lme_cls = 'nt'
            elif lme_basis < 0 and lme03_trend is not None and lme03_trend < 0:
                lme_conclusion = 'Contango扩大，海外结构继续走弱'; lme_status = '偏空'; lme_cls = 'down'
            elif lme_basis < 0 and lme03_trend is not None and lme03_trend > 0:
                lme_conclusion = 'Contango收窄，海外结构边际改善，仍需确认'; lme_status = '关注'; lme_cls = 'nt'
            if domestic_premium < 0 and domestic_trend is not None and domestic_trend < 0:
                domestic_conclusion = '贴水扩大，国内结构正在走弱'
                domestic_status, domestic_cls = '偏空', 'down'
            elif domestic_premium < 0 and domestic_trend is not None and domestic_trend > 0:
                domestic_conclusion = '贴水收窄，国内结构边际改善，仍需确认'
                domestic_status, domestic_cls = '关注', 'nt'
            elif domestic_premium < 0:
                domestic_conclusion = '现货贴水，国内结构偏弱'
                domestic_status, domestic_cls = '偏空', 'down'
            elif domestic_premium > 0 and domestic_trend is not None and domestic_trend < 0:
                domestic_conclusion = '升水走弱，国内结构边际转弱，跟踪是否转贴水'
                domestic_status, domestic_cls = '关注', 'nt'
            elif domestic_premium > 0:
                domestic_conclusion = '现货升水，国内结构仍有支撑'
                domestic_status, domestic_cls = '偏多', 'up'
            else:
                domestic_conclusion = '国内现货结构持平'
                domestic_status, domestic_cls = '中性', 'nt'
            if social_trend is not None:
                social_conclusion = '国内社库20日 %s，%s' % (shared.pct(social_trend), '去库支撑近月' if social_trend < 0 else '累库压制近月')
                social_status = '去库偏多' if social_trend < 0 else '累库偏空'
                social_cls = 'up' if social_trend < 0 else 'down'
            else:
                social_conclusion = '国内社库趋势暂不可得，与海外仓单共同验证'
                social_status, social_cls = '跟踪', 'nt'
            supply_label = '供应·矿端' if metal == 'CU' else '供应·原生'
            demand_label = '需求·加工' if metal == 'CU' else '需求·电池'
            supply_status = '偏多' if marginal_scores['供应'] > 0 else ('偏空' if marginal_scores['供应'] < 0 else '跟踪')
            demand_status = '偏多' if marginal_scores['需求'] > 0 else ('偏空' if marginal_scores['需求'] < 0 else '跟踪')
            rows = [
                ('价格', '沪%s %s 元/吨（%s）｜LME%s %s 美元/吨（%s）' % (cn, shared.fmt(latest['shfe'][1]), latest['shfe'][0], cn, shared.fmt(latest['lme'][1]), latest['lme'][0]), '近5个交易日 %s，价格%s' % (shared.pct(chg5), price_word), price_word, _cls(chg5)),
                ('结构·LME', 'LME 0-3 %s 美元/吨（%s，20日 %s）' % (shared.fmt(latest['lme03'][1]), latest['lme03'][0], shared.pct(lme03_trend)), lme_conclusion, lme_status, lme_cls),
                ('结构·国内', '国内现货%s %s 元/吨（%s，20日 %s）' % (domestic_structure, shared.fmt(latest['domestic'][1]), latest['domestic'][0], shared.pct(domestic_trend)), domestic_conclusion, domestic_status, domestic_cls),
                ('海外库存', 'LME%s库存 %s 吨（%s）' % (cn, shared.fmt(last(charts[lme_key])[1]), last(charts[lme_key])[0]), '近5个交易日 %s，海外%s' % (shared.pct(stock_chg), stock_word), '%s%s' % (stock_word, '偏多' if stock_word == '去库' else '偏空'), 'up' if stock_word == '去库' else 'dn'),
                ('国内库存', '%s社库 %s 万吨（%s）' % (cn, shared.fmt(latest['social_stock'][1]), latest['social_stock'][0]), social_conclusion, social_status, social_cls),
                (supply_label, '%s %s %s（%s）' % ('铜精矿TC' if metal == 'CU' else '原生铅产量', shared.fmt(latest['tc'][1]) if metal == 'CU' else shared.fmt(output[1] if output else None), '美元/吨' if metal == 'CU' else '万吨', latest['tc'][0] if metal == 'CU' else (output[0] if output else '—')), '供应方向按边际块证据确认，不单独外推', supply_status, _cls(marginal_scores['供应'])),
                (demand_label, '%s %s%s（%s）' % ('精铜制杆产能利用率' if metal == 'CU' else '铅酸蓄电池开工率', shared.fmt(demand[1]) if demand else '—', '%' if demand else '', demand[0] if demand else '—'), '需求方向按开工/利用率边际变化确认，不单独外推', demand_status, _cls(marginal_scores['需求'])),
                ('策略观察', '价格、LME结构、国内结构、库存、供应和需求交叉验证', '信号未同向前不作单边外推；铅重点盯国内升贴水是否转弱，铜重点盯TC与库存', '跟踪', 'nt'),
            ]
            for label, datum, conclusion, status, cls in rows:
                parts.append('<tr><td class="c">%s</td><td>%s</td><td>%s</td><td><span class="dir %s">%s</span></td></tr>' % (
                    shared.esc(label), shared.esc(datum), shared.esc(conclusion), cls, shared.esc(status)))
            parts.append('</table></div>')
            parts.append(marginal_html)
            parts.append('<footer>数据更新至 %s ｜ %s研究周报 ｜ ⚙️自动构建，仅供研究参考。</footer>' % (shared.esc(asof), cn))
        elif sid == 'framework':
            if metal == 'CU':
                text = '矿端线：铜精矿 → TC/进口 → 冶炼与硫酸收益|锭端线：精炼铜 → LME/SHFE/社库 → 升贴水与进口盈亏|需求线：铜杆/铜管/铜板带 → 电线电缆、电网、家电与新能源'
            else:
                text = '矿端线：铅精矿 → TC/进口 → 原生铅|再生线：废电瓶 → 再生粗铅 → 精炼铅|锭端线：精炼铅 → LME/SHFE/社库 → 现货结构|需求线：铅酸蓄电池 → 汽车启动、储能与替换需求'
            blocks = ''.join('<div class="lead">%s</div><p>按来源、产量、库存、利润和终端开工交叉验证，不以单一指标代替全产业链。</p>' % shared.esc(x) for x in text.split('|'))
            parts.append('<details class="fold" id="framework" open><summary>%s产业研究框架</summary><div class="foldbody">%s<p><a href="https://amiya866.github.io/metals-framework/">打开金属研究框架 ↗</a> · <a href="https://amiya866.github.io/caibao-hub/">打开%s企财报 ↗</a></p></div></details>' % (cn, blocks, cn))
        elif sid == 'cost':
            parts.append('<div class="lead">成本、加工费与财报项目</div>')
            parts.append(shared.chart_group(charts, keys))
            wb = load_workbook(cfg['financial'], read_only=True, data_only=True)
            for sheet_name, limit in cfg['sheets']:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                headers = list(rows[0])
                period = '26Q3' if '26Q3' in headers else '26Q2'
                idx = headers.index(period)
                guide_idx = headers.index('FY2026指引')
                parts.append('<div class="lead">%s · %s</div><div class="tablewrap"><table class="bal"><tr><th>公司</th><th>国家/项目</th><th>%s</th><th>FY2026指引</th></tr>' % (sheet_name, period, period))
                for row in rows[1:limit + 1]:
                    if row[idx] in (None, '-'):
                        continue
                    parts.append('<tr><td>%s</td><td>%s</td><td class="num">%s</td><td>%s</td></tr>' % (
                        shared.esc(row[0]), shared.esc(row[2] or row[1]), shared.esc(row[idx]), shared.esc(row[guide_idx] or '—')))
                parts.append('</table></div>')
            wb.close()
        elif sid == 'technical':
            parts.append('<div class="tech-grid">')
            cards = [
                ('日线趋势', '偏多' if current_rsi and current_rsi > 55 else ('偏空' if current_rsi and current_rsi < 45 else '震荡'),
                 '收盘 %s；MA5 %s / MA20 %s；RSI14 %.1f；价格%s。' % (shared.fmt(shfe_values[-1]), shared.fmt(ma5), shared.fmt(ma20), current_rsi or 0, price_word),
                 'up' if ma5 > ma20 else 'down'),
                ('MACD动量', '偏多' if macd[-1] > signal[-1] else '偏空',
                 'DIF %s / DEA %s；柱体 %s。' % (shared.fmt(macd[-1]), shared.fmt(signal[-1]), shared.fmt(macd[-1] - signal[-1])),
                 'up' if macd[-1] > signal[-1] else 'down'),
                ('关键区间', '观察', '近20日低 %s / 高 %s；结合库存和升贴水验证，不以单一技术位直接下单。' % (shared.fmt(min(shfe_values[-20:])), shared.fmt(max(shfe_values[-20:]))), 'neutral'),
            ]
            for frame, status, detail, tone in cards:
                parts.append('<article class="tech-card %s"><h3>%s<span class="score">%s</span></h3><p>%s</p></article>' % (tone, frame, status, shared.esc(detail)))
            parts.append('</div><div class="tech-note">技术分析由最新沪%s日线自动计算，仅供研究参考。</div><div class="lead">沪%s主连 · 日线</div>' % (cn, cn))
            parts.append(shared.chart_group(charts, keys))
        elif sid == 'sources':
            summary = freshness['summary']
            parts.append('<details class="fold" id="datagap" open><summary>数据口径、缺口与新鲜度（%d项全量审计）</summary><div class="foldbody">' % summary['total'])
            parts.append('<div class="fresh-grid">')
            for code, label in [('fresh', '正常'), ('watch', '关注'), ('stale', '滞后'), ('missing', '缺失'), ('review', '待核')]:
                parts.append('<button type="button" class="fresh-kpi %s" data-filter="%s"><span>%s</span><strong>%d</strong><small>点击筛选</small></button>' % (code, code, label, summary[code]))
            parts.append('<button type="button" class="fresh-kpi total" data-filter="all"><span>全部指标</span><strong>%d</strong><small>重置筛选</small></button></div>' % summary['total'])
            parts.append('<p>来源：数据缓存_v1.db、guan/新浪行情缓存、%s企财报真源。日度 7/14 天、周度 10/21 天、月度 45/75 天、季度 150/240 天为正常/关注阈值；缺失和滞后不补零，不以旧值冒充最新。</p></div></details>' % cn)
            parts.append('<div class="fresh-legend">当前显示 <b id="freshCount">%d</b> / %d 项｜点击上方卡片筛选</div>' % (summary['total'], summary['total']))
            parts.append('<div class="tablewrap"><table class="bal fresh-table"><tr><th>模块/指标</th><th>来源</th><th>频率</th><th>单位</th><th>最新标签/值</th><th>新鲜度</th><th>处理</th></tr>')
            for row in freshness['rows']:
                age = '，%d天' % row['age_days'] if row['age_days'] is not None else ''
                parts.append('<tr data-fresh="%s"><td><b>%s</b><br><span class="src-note">%s｜%s</span></td><td>%s</td><td>%s</td><td>%s</td><td>%s<br><span class="src-note">%s%s</span></td><td><span class="fresh-badge %s">%s</span></td><td>%s</td></tr>' % (
                    row['status'], shared.esc(row['name']), shared.esc(row['module']), shared.esc(row['key']), shared.esc(row['source']),
                    shared.esc(row['frequency']), shared.esc(row['unit']), shared.esc(row['latest']), shared.esc(row['latest_value']), shared.esc(age),
                    row['status'], shared.esc(row['status_label']), shared.esc(row['action'])))
            parts.append('</table></div>')
        else:
            parts.append('<div class="lead">%s · 最新读数逐项标注日期</div>' % shared.esc(title))
            parts.append(shared.chart_group(charts, keys))
        parts.append('</div></section>')

    payload = json.dumps(charts, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    other = {'SN': 'tin-weekly', 'NI': 'nickel-weekly', 'ZN': 'zinc-weekly', 'AL': 'aluminum-weekly'}
    switches = ''.join('<a%s href="%s">%s</a>' % (' class="active"' if code == metal else '',
                       '#sumsec' if code == metal else '../%s/#sumsec' % (CONFIG[code]['out'] if code in CONFIG else other[code]),
                       CONFIG[code]['cn'] if code in CONFIG else {'SN': '锡', 'NI': '镍', 'ZN': '锌', 'AL': '铝'}[code])
                       for code in ('SN', 'NI', 'ZN', 'AL', 'CU', 'PB'))
    js_config = json.dumps({'cn': cn, 'prefix': cfg['prefix'], 'supply': cfg['js_supply'], 'demand': cfg['js_demand']}, ensure_ascii=False)
    doc = '''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow,noarchive"><link href="data:," rel="icon">
<title>研究周报 · __CN__</title><style>__BASE_CSS__</style><link href="weekly_mobile.css" rel="stylesheet"><link href="weekly_polish.css" rel="stylesheet"><link href="research.css" rel="stylesheet"><link href="../weekly-light.css" rel="stylesheet"><style>.chart-box{height:240px;position:relative}.chart-box canvas{display:block;width:100%;height:100%}.tablewrap{overflow-x:auto}.page-body p{font-size:14.5px}.cbox .stats{line-height:1.55}.fresh-grid{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:9px;margin:12px 0}.fresh-kpi{border:1px solid var(--line);background:var(--cardbg);padding:10px;text-align:left;cursor:pointer}.fresh-kpi span,.fresh-kpi small{display:block;color:var(--faint);font-size:11px}.fresh-kpi strong{display:block;font-size:25px}.fresh-kpi.fresh strong{color:#188865}.fresh-kpi.watch strong{color:#a3760e}.fresh-kpi.stale strong,.fresh-kpi.missing strong{color:#c44b5d}.fresh-kpi.review strong{color:#7c4dff}.fresh-kpi.active{outline:2px solid var(--blue);outline-offset:1px}.fresh-legend{margin:12px 0;color:var(--grey);font-size:13px}.fresh-badge{display:inline-block;padding:1px 7px;border:1px solid var(--line);white-space:nowrap}.fresh-badge.fresh{color:#188865}.fresh-badge.watch{color:#a3760e}.fresh-badge.stale,.fresh-badge.missing{color:#c44b5d}.fresh-badge.review{color:#7c4dff}.src-note{color:var(--faint);font-size:11px}.fresh-table{min-width:1080px}.fresh-table td{vertical-align:top}@media(max-width:900px){.fresh-grid{grid-template-columns:repeat(3,1fr)}}</style></head>
<body><header class="research-topbar"><div class="research-brand"><a aria-label="返回整合研究站" class="back-home" href="../index.html">← 整合研究站</a><div class="brand-copy"><p class="research-eyebrow">COMMODITY WEEKLY RESEARCH · __EN__</p><h1>研究周报 <span class="commodity-current">__CN__</span></h1><p>__SUBTITLE__</p></div></div><div class="research-actions"><nav class="commodity-switch" aria-label="品种切换">__SWITCHES__</nav><a class="back-home" href="https://amiya866.github.io/caibao-hub/" target="_blank" rel="noopener">__CN__企财报 ↗</a><span class="research-stamp">数据截至 <b>__ASOF__</b></span><button aria-pressed="false" class="view-toggle" id="viewToggle" type="button">切换手机版</button></div></header>
<div class="research-workspace"><aside class="research-sidebar top-navigation" id="researchSidebar"><button aria-controls="treeNav" aria-expanded="false" class="menu-toggle" id="menuToggle" type="button">研究目录 <span>＋</span></button><div class="sidebar-heading"><span>RESEARCH INDEX</span><b>研究目录</b></div><nav aria-label="__CN__研究分层目录" class="tree-nav top-nav-inner" id="treeNav"></nav></aside><div class="research-content"><div aria-live="polite" class="breadcrumb" id="breadcrumb"></div><div class="wrap" id="reportWrap">__SECTIONS__</div><section aria-live="polite" class="tool-page" hidden id="toolPage"><div class="tool-head"><div><p class="research-eyebrow">RESEARCH WORKBENCH</p><h2 id="toolTitle">研究工作台</h2></div><button id="closeTool" type="button">← 返回周报</button></div><iframe hidden id="toolFrame" title="__CN__研究工作台"></iframe></section></div></div>
<script src="lib/chart.min.js"></script><script>const charts=__PAYLOAD__;document.querySelectorAll('canvas[data-chart]').forEach(function(canvas){const c=charts[canvas.dataset.chart];if(!c)return;const ds=(c.datasets||[]).map(function(d){return Object.assign({borderWidth:1.5,pointRadius:0,spanGaps:true},d);});const cfg={type:'line',data:{labels:c.labels||[],datasets:ds},options:{responsive:true,maintainAspectRatio:false,animation:false,interaction:{mode:'index',intersect:false},plugins:{legend:{position:'bottom',labels:{boxWidth:12,font:{family:'KaiTi,"楷体",serif',size:10}}}},scales:{x:{ticks:{maxTicksLimit:8,font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}},y:{ticks:{font:{family:'KaiTi,"楷体",serif',size:10}},grid:{color:'#e3e3e3'}}}}};new Chart(canvas,cfg);});</script><script src="weekly_mobile.js"></script><script>window.WEEKLY_CONFIG=__JS_CONFIG__;</script><script src="research.js"></script><script>document.querySelectorAll('.fresh-kpi').forEach(function(btn){btn.addEventListener('click',function(){const code=btn.dataset.filter;let shown=0;document.querySelectorAll('.fresh-kpi').forEach(function(x){x.classList.toggle('active',x===btn)});document.querySelectorAll('tr[data-fresh]').forEach(function(row){const visible=code==='all'||row.dataset.fresh===code;row.hidden=!visible;if(visible)shown++});document.getElementById('freshCount').textContent=shown})});</script></body></html>'''
    doc = (doc.replace('__BASE_CSS__', base_css).replace('__ASOF__', shared.esc(asof))
           .replace('__SECTIONS__', ''.join(parts)).replace('__PAYLOAD__', payload)
           .replace('__JS_CONFIG__', js_config).replace('__SWITCHES__', switches)
           .replace('__CN__', cn).replace('__EN__', cfg['en']).replace('__SUBTITLE__', cfg['subtitle']))
    tmp = out / 'index.html.tmp'
    tmp.write_text(doc, encoding='utf-8')
    os.replace(tmp, out / 'index.html')
    summary = {'title': '研究周报 · %s' % cn, 'asof': asof, 'headline': headline, 'metrics': metrics,
               'local_url': '%s/#sumsec' % cfg['out'], 'source_url': 'https://amiya866.github.io/caibao-hub/',
               'charts': len(charts), 'freshness': freshness['summary'], 'sync_status': 'ok'}
    (BASE / 'data' / ('%s_weekly.json' % cfg['prefix'])).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    update_daily(cfg, summary)
    print(json.dumps(summary, ensure_ascii=False))
